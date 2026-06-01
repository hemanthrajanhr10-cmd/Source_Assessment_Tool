"""
Snowflake Assessment Service.

Orchestrates a full head-to-toe Snowflake platform analysis:
  1.  Verify connection + account metadata
  2.  Enumerate warehouses (compute layer)
  3.  Enumerate databases
  4.  Enumerate schemas (across all accessible databases)
  5.  Enumerate tables & views (INFORMATION_SCHEMA per database)
  6.  Inventory platform objects (stages, pipes, tasks, streams, procedures, functions, ...)
  7.  Enumerate users + roles (RBAC posture)
  8.  Security posture (network policies, masking/row-access policies, MFA coverage)
  9.  Query performance metrics (ACCOUNT_USAGE.QUERY_HISTORY — 7 days)
  10. Storage usage (ACCOUNT_USAGE.STORAGE_USAGE)
  11. Cost/credit metering (ACCOUNT_USAGE.WAREHOUSE_METERING_HISTORY — 30 days)
  12. Generate Excel report
  13. Generate Word report
"""

import io
import uuid
from datetime import datetime, timezone
from typing import Optional

from app.core.logging import get_logger
from app.db import snowflake_client as client
from app.models.snowflake_requests import (
    SnowflakeAssessmentRequest,
    SnowflakeAssessmentResult,
    SnowflakeAccountInfo,
    SnowflakeWarehouse,
    SnowflakeWarehouseMetrics,
    SnowflakeDatabase,
    SnowflakeSchema,
    SnowflakeTable,
    SnowflakeDatabaseSummary,
    SnowflakeObjectInventory,
    SnowflakeUserProfile,
    SnowflakeSecurityPosture,
    SnowflakeQueryMetrics,
    SnowflakeStorageMetrics,
    SnowflakeCostMetrics,
)

logger = get_logger(__name__)

# ── In-memory job store ───────────────────────────────────────────────────────

_jobs: dict[str, dict] = {}


def create_job(request: SnowflakeAssessmentRequest, account: str) -> str:
    job_id = str(uuid.uuid4())
    _jobs[job_id] = {
        "job_id": job_id,
        "label": request.label,
        "account": account,
        "status": "pending",
        "progress_message": None,
        "error": None,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "completed_at": None,
        "results": None,
        "excel_bytes": None,
        "word_bytes": None,
    }
    return job_id


def get_job(job_id: str) -> Optional[dict]:
    return _jobs.get(job_id)


def list_jobs() -> list[dict]:
    return sorted(_jobs.values(), key=lambda j: j["created_at"], reverse=True)


def _update(job_id: str, **kwargs) -> None:
    if job_id in _jobs:
        _jobs[job_id].update(kwargs)


# ── Assessment steps (used for frontend progress terminal) ────────────────────

STEPS = [
    "Verifying account connection",
    "Fetching account metadata",
    "Enumerating warehouses",
    "Enumerating databases",
    "Enumerating schemas",
    "Enumerating tables & views",
    "Inventorying platform objects",
    "Enumerating users",
    "Enumerating roles",
    "Assessing security posture",
    "Fetching query performance metrics",
    "Fetching storage usage",
    "Fetching warehouse credit metering",
    "Generating Excel report",
    "Generating Word report",
]


# ── Assessment runner ─────────────────────────────────────────────────────────

def run_assessment(job_id: str, request: SnowflakeAssessmentRequest) -> None:
    """Full Snowflake assessment — runs in a FastAPI background task."""

    def _step(msg: str) -> None:
        logger.info("[snowflake:%s] %s", job_id[:8], msg)
        _update(job_id, status="running", progress_message=msg)

    try:
        # ── Step 1: Get connection from auth session ───────────────────────────
        _step(STEPS[0])
        conn = client.get_auth_connection(request.auth_id)
        if conn is None:
            raise RuntimeError(
                f"Auth session {request.auth_id!r} is not authenticated. "
                "Complete browser login before starting assessment."
            )

        # ── Step 2: Account metadata ──────────────────────────────────────────
        _step(STEPS[1])
        raw_info = client.get_account_info(conn)
        account_info = SnowflakeAccountInfo(
            account_name=raw_info.get("account_name") or raw_info.get("account_locator") or "Unknown",
            organization_name=raw_info.get("organization_name"),
            account_locator=raw_info.get("account_locator"),
            region=raw_info.get("region"),
            snowflake_version=raw_info.get("snowflake_version"),
            current_role=raw_info.get("current_role"),
            current_warehouse=raw_info.get("current_warehouse"),
            current_user=raw_info.get("current_user"),
        )

        # ── Step 3: Warehouses ────────────────────────────────────────────────
        _step(STEPS[2])
        raw_wh = client.list_warehouses(conn)
        warehouses = [SnowflakeWarehouse(**w) for w in raw_wh]
        wh_by_size: dict[str, int] = {}
        for w in warehouses:
            s = w.size or "Unknown"
            wh_by_size[s] = wh_by_size.get(s, 0) + 1

        warehouse_metrics = SnowflakeWarehouseMetrics(
            total_warehouses=len(warehouses),
            active_warehouses=sum(1 for w in warehouses if w.state.upper() == "STARTED"),
            suspended_warehouses=sum(1 for w in warehouses if w.state.upper() == "SUSPENDED"),
            warehouses_by_size=wh_by_size,
            multi_cluster_warehouses=sum(1 for w in warehouses if w.max_cluster_count and w.max_cluster_count > 1),
        )

        # ── Step 4: Databases ─────────────────────────────────────────────────
        _step(STEPS[3])
        raw_dbs = client.list_databases(conn)
        databases = [SnowflakeDatabase(**d) for d in raw_dbs]
        db_names = [d.name for d in databases[:request.max_databases]]

        # ── Step 5: Schemas ───────────────────────────────────────────────────
        _step(STEPS[4])
        raw_schemas: list[dict] = []
        try:
            raw_schemas = client.list_schemas(conn, db_names)
        except Exception as exc:
            logger.warning("list_schemas partial failure: %s", exc)
        schemas = [SnowflakeSchema(**s) for s in raw_schemas]

        # ── Step 6: Tables & Views ────────────────────────────────────────────
        _step(STEPS[5])
        all_tables: list[SnowflakeTable] = []
        total_bytes = 0
        for db_name in db_names:
            try:
                raw_tables = client.list_tables_in_db(conn, db_name)
                for t in raw_tables:
                    total_bytes += int(t.get("bytes") or 0)
                    all_tables.append(SnowflakeTable(**t))
            except Exception as exc:
                logger.warning("list_tables db=%s: %s", db_name, exc)

        db_summary = SnowflakeDatabaseSummary(
            total_databases=len(databases),
            total_schemas=len(schemas),
            total_tables=sum(1 for t in all_tables if t.table_type == "BASE TABLE"),
            total_views=sum(1 for t in all_tables if t.table_type == "VIEW"),
            total_external_tables=sum(1 for t in all_tables if "EXTERNAL" in t.table_type.upper()),
            total_materialized_views=sum(1 for t in all_tables if "MATERIALIZED" in t.table_type.upper()),
            total_size_bytes=total_bytes,
        )

        # ── Step 7: Platform object inventory ────────────────────────────────
        _step(STEPS[6])
        shares = {"outbound": 0, "inbound": 0}
        obj_counts = {}
        _show_map = {
            "stages": "SHOW STAGES IN ACCOUNT",
            "pipes": "SHOW PIPES IN ACCOUNT",
            "tasks": "SHOW TASKS IN ACCOUNT",
            "streams": "SHOW STREAMS IN ACCOUNT",
            "procedures": "SHOW PROCEDURES IN ACCOUNT",
            "functions": "SHOW USER FUNCTIONS IN ACCOUNT",
            "sequences": "SHOW SEQUENCES IN ACCOUNT",
            "file_formats": "SHOW FILE FORMATS IN ACCOUNT",
            "resource_monitors": "SHOW RESOURCE MONITORS",
            "network_policies": "SHOW NETWORK POLICIES",
        }
        for key, sql in _show_map.items():
            try:
                obj_counts[key] = client.count_objects(conn, sql)
            except Exception as exc:
                logger.warning("count %s: %s", key, exc)
                obj_counts[key] = 0

        # Dynamic tables (Snowflake 2023+)
        dynamic_tables = 0
        try:
            dynamic_tables = client.count_objects(conn, "SHOW DYNAMIC TABLES IN ACCOUNT")
        except Exception:
            pass

        try:
            shares = client.get_shares(conn)
        except Exception as exc:
            logger.warning("get_shares: %s", exc)

        masking_policies = 0
        row_access_policies = 0
        try:
            masking_policies = client.count_masking_policies(conn)
        except Exception:
            pass
        try:
            row_access_policies = client.count_row_access_policies(conn)
        except Exception:
            pass

        obj_inventory = SnowflakeObjectInventory(
            stages=obj_counts.get("stages", 0),
            pipes=obj_counts.get("pipes", 0),
            tasks=obj_counts.get("tasks", 0),
            streams=obj_counts.get("streams", 0),
            procedures=obj_counts.get("procedures", 0),
            functions=obj_counts.get("functions", 0),
            sequences=obj_counts.get("sequences", 0),
            file_formats=obj_counts.get("file_formats", 0),
            dynamic_tables=dynamic_tables,
            shares_outbound=shares["outbound"],
            shares_inbound=shares["inbound"],
            resource_monitors=obj_counts.get("resource_monitors", 0),
            network_policies=obj_counts.get("network_policies", 0),
            masking_policies=masking_policies,
            row_access_policies=row_access_policies,
        )

        # ── Step 8: Users ─────────────────────────────────────────────────────
        _step(STEPS[7])
        raw_users: list[dict] = []
        try:
            raw_users = client.list_users(conn)
        except Exception as exc:
            logger.warning("list_users: %s", exc)

        disabled_users = sum(1 for u in raw_users if u.get("disabled"))
        mfa_missing = sum(1 for u in raw_users if not u.get("has_mfa") and not u.get("disabled"))
        admin_users = sum(1 for u in raw_users if "ADMIN" in str(u.get("default_role") or "").upper() or "SYSADMIN" in str(u.get("default_role") or "").upper())

        # ── Step 9: Roles ─────────────────────────────────────────────────────
        _step(STEPS[8])
        raw_roles: list[dict] = []
        try:
            raw_roles = client.list_roles(conn)
        except Exception as exc:
            logger.warning("list_roles: %s", exc)

        _system_roles = {"ACCOUNTADMIN", "SYSADMIN", "SECURITYADMIN", "USERADMIN", "ORGADMIN", "PUBLIC"}
        custom_roles = sum(1 for r in raw_roles if r.get("name", "").upper() not in _system_roles)

        user_profile = SnowflakeUserProfile(
            total_users=len(raw_users),
            disabled_users=disabled_users,
            users_without_mfa=mfa_missing,
            admin_users=admin_users,
            service_accounts=sum(1 for u in raw_users if "SVC" in u.get("name", "").upper() or "SERVICE" in u.get("name", "").upper()),
            total_roles=len(raw_roles),
            custom_roles=custom_roles,
            system_roles=len(raw_roles) - custom_roles,
        )

        # ── Step 10: Security posture ─────────────────────────────────────────
        _step(STEPS[9])
        security_posture = SnowflakeSecurityPosture(
            network_policies_count=obj_counts.get("network_policies", 0),
            users_without_mfa=mfa_missing,
            users_with_default_role_public=sum(1 for u in raw_users if u.get("default_role", "").upper() == "PUBLIC"),
            masking_policies_count=masking_policies,
            row_access_policies_count=row_access_policies,
            shares_total=shares["outbound"] + shares["inbound"],
            resource_monitors_count=obj_counts.get("resource_monitors", 0),
        )

        # ── Step 11: Query metrics ────────────────────────────────────────────
        query_metrics = SnowflakeQueryMetrics()
        if request.include_query_history:
            _step(STEPS[10])
            try:
                raw_qm = client.get_query_history(conn)
                query_metrics = SnowflakeQueryMetrics(
                    total_queries_last_7d=raw_qm.get("total_queries", 0),
                    failed_queries_last_7d=raw_qm.get("failed_queries", 0),
                    avg_execution_ms=raw_qm.get("avg_execution_ms", 0.0),
                    p95_execution_ms=raw_qm.get("p95_execution_ms", 0.0),
                    bytes_scanned_total=raw_qm.get("bytes_scanned", 0),
                    bytes_spilled_local=raw_qm.get("bytes_spilled_local", 0),
                    bytes_spilled_remote=raw_qm.get("bytes_spilled_remote", 0),
                    most_expensive_queries=raw_qm.get("top_expensive", []),
                    query_error_types=raw_qm.get("error_types", {}),
                )
            except Exception as exc:
                logger.warning("query_history failed (non-fatal): %s", exc)

        # ── Step 12: Storage ──────────────────────────────────────────────────
        storage_metrics = SnowflakeStorageMetrics()
        if request.include_storage_usage:
            _step(STEPS[11])
            try:
                raw_st = client.get_storage_usage(conn)
                storage_metrics = SnowflakeStorageMetrics(
                    storage_bytes=raw_st.get("storage_bytes", 0),
                    stage_bytes=raw_st.get("stage_bytes", 0),
                    failsafe_bytes=raw_st.get("failsafe_bytes", 0),
                    total_bytes=sum(raw_st.get(k, 0) for k in ("storage_bytes", "stage_bytes", "failsafe_bytes")),
                )
            except Exception as exc:
                logger.warning("storage_usage failed (non-fatal): %s", exc)

        # ── Step 13: Credits ──────────────────────────────────────────────────
        cost_metrics = SnowflakeCostMetrics()
        if request.include_warehouse_metering:
            _step(STEPS[12])
            try:
                raw_cr = client.get_warehouse_credits(conn)
                cost_metrics = SnowflakeCostMetrics(
                    credits_used_last_30d=raw_cr.get("total_credits", 0.0),
                    compute_credits=raw_cr.get("compute_credits", 0.0),
                    cloud_services_credits=raw_cr.get("cloud_services_credits", 0.0),
                    top_warehouses_by_credit=raw_cr.get("top_wh", []),
                )
            except Exception as exc:
                logger.warning("warehouse_credits failed (non-fatal): %s", exc)

        # ── Assemble result ───────────────────────────────────────────────────
        result = SnowflakeAssessmentResult(
            job_id=job_id,
            label=request.label,
            assessed_at=datetime.now(timezone.utc).isoformat(),
            status="completed",
            account_info=account_info,
            warehouse_metrics=warehouse_metrics,
            database_summary=db_summary,
            object_inventory=obj_inventory,
            user_profile=user_profile,
            security_posture=security_posture,
            query_metrics=query_metrics,
            storage_metrics=storage_metrics,
            cost_metrics=cost_metrics,
            warehouses=warehouses[:100],
            databases=databases[:200],
            schemas=schemas[:500],
            tables=all_tables[:1000],
            users=raw_users[:500],
            roles=raw_roles[:200],
        )

        # ── Step 14 & 15: Reports ──────────────────────────────────────────────
        _step(STEPS[13])
        excel_bytes = _build_excel(result)

        _step(STEPS[14])
        word_bytes = _build_word(result)

        _update(
            job_id,
            status="completed",
            progress_message="Assessment complete",
            completed_at=datetime.now(timezone.utc).isoformat(),
            results=result.model_dump(),
            excel_bytes=excel_bytes,
            word_bytes=word_bytes,
        )
        logger.info("[snowflake:%s] Assessment completed successfully", job_id[:8])

    except Exception as exc:
        logger.exception("[snowflake:%s] Assessment failed: %s", job_id[:8], exc)
        _update(
            job_id,
            status="failed",
            error=str(exc),
            completed_at=datetime.now(timezone.utc).isoformat(),
            progress_message=f"Failed: {exc}",
        )


# ── Excel report ──────────────────────────────────────────────────────────────

SNOW_CYAN  = "00B8E6"
SNOW_NAVY  = "0A1628"
SNOW_TEAL  = "0099CC"
SNOW_LIGHT = "EBF4FF"
SNOW_PALE  = "F0F8FF"
WHITE      = "FFFFFF"
DARK_GRAY  = "2C3E50"
LIGHT_GRAY = "F5F5F5"
FONT_NAME  = "Calibri"


def _build_excel(result: SnowflakeAssessmentResult) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    def _font(bold=False, size=11, color=DARK_GRAY, italic=False):
        return Font(name=FONT_NAME, bold=bold, size=size, color=color, italic=italic)

    def _fill(c: str):
        return PatternFill("solid", fgColor=c)

    def _border():
        s = Side(style="thin")
        return Border(left=s, right=s, top=s, bottom=s)

    def _align(h="left", v="center", wrap=False):
        return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

    def _hdr(ws, row, col, val, bg=SNOW_NAVY, fg=WHITE):
        c = ws.cell(row=row, column=col, value=val)
        c.font = _font(bold=True, color=fg)
        c.fill = _fill(bg)
        c.alignment = _align("center")
        c.border = _border()

    def _cell(ws, row, col, val, shade=False):
        c = ws.cell(row=row, column=col, value=val)
        c.font = _font()
        c.fill = _fill(SNOW_PALE if shade else LIGHT_GRAY)
        c.alignment = _align()
        c.border = _border()

    def _auto_width(ws):
        for col in ws.columns:
            col_ltr = get_column_letter(col[0].column)
            w = max((len(str(cell.value)) for cell in col if cell.value), default=0)
            ws.column_dimensions[col_ltr].width = min(w + 4, 50)

    wb = Workbook()
    wb.remove(wb.active)

    # ── Overview ──────────────────────────────────────────────────────────────
    ws = wb.create_sheet("Overview")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    t = ws["A1"]
    acct = result.account_info.account_name if result.account_info else (result.label or "Snowflake")
    t.value = f"Snowflake Assessment — {acct}"
    t.font = Font(name=FONT_NAME, bold=True, size=16, color=WHITE)
    t.fill = _fill(SNOW_NAVY)
    t.alignment = _align("center")

    ws.merge_cells("A2:D2")
    sub = ws["A2"]
    sub.value = f"Assessed: {result.assessed_at[:19].replace('T', ' ')} UTC  |  Job: {result.job_id[:8]}"
    sub.font = _font(italic=True, size=10, color=SNOW_CYAN)
    sub.fill = _fill(SNOW_LIGHT)
    sub.alignment = _align("center")

    overview_rows: list[tuple] = []
    if result.account_info:
        ai = result.account_info
        overview_rows += [
            ("Account Name", ai.account_name or ""),
            ("Account Locator", ai.account_locator or ""),
            ("Organization", ai.organization_name or ""),
            ("Region", ai.region or ""),
            ("Snowflake Version", ai.snowflake_version or ""),
            ("Current Role", ai.current_role or ""),
            ("Current User", ai.current_user or ""),
        ]
    if result.warehouse_metrics:
        wm = result.warehouse_metrics
        overview_rows += [
            ("Total Warehouses", wm.total_warehouses),
            ("Active Warehouses", wm.active_warehouses),
            ("Suspended Warehouses", wm.suspended_warehouses),
            ("Multi-Cluster Warehouses", wm.multi_cluster_warehouses),
        ]
    if result.database_summary:
        ds = result.database_summary
        overview_rows += [
            ("Total Databases", ds.total_databases),
            ("Total Schemas", ds.total_schemas),
            ("Total Tables", ds.total_tables),
            ("Total Views", ds.total_views),
            ("Total External Tables", ds.total_external_tables),
            ("Total Materialized Views", ds.total_materialized_views),
            ("Total Data Size (bytes)", ds.total_size_bytes),
        ]
    if result.user_profile:
        up = result.user_profile
        overview_rows += [
            ("Total Users", up.total_users),
            ("Disabled Users", up.disabled_users),
            ("Users Without MFA", up.users_without_mfa),
            ("Admin Users", up.admin_users),
            ("Total Roles", up.total_roles),
            ("Custom Roles", up.custom_roles),
        ]

    _hdr(ws, 4, 1, "Metric")
    _hdr(ws, 4, 2, "Value")
    for i, (k, v) in enumerate(overview_rows, start=5):
        _cell(ws, i, 1, k, i % 2 == 0)
        _cell(ws, i, 2, v, i % 2 == 0)
    _auto_width(ws)

    # ── Warehouses ────────────────────────────────────────────────────────────
    if result.warehouses:
        ws2 = wb.create_sheet("Warehouses")
        ws2.sheet_view.showGridLines = False
        headers = ["Name", "State", "Type", "Size", "Auto-Suspend (s)", "Auto-Resume",
                   "Max Clusters", "Running", "Queued", "Owner"]
        for ci, h in enumerate(headers, 1):
            _hdr(ws2, 1, ci, h)
        for i, w in enumerate(result.warehouses, start=2):
            shade = i % 2 == 0
            for ci, v in enumerate([
                w.name, w.state, w.wh_type, w.size, w.auto_suspend,
                "Yes" if w.auto_resume else "No",
                w.max_cluster_count or 1, w.running, w.queued, w.owner or "",
            ], 1):
                _cell(ws2, i, ci, v, shade)
        ws2.auto_filter.ref = f"A1:J{len(result.warehouses)+1}"
        ws2.freeze_panes = "A2"
        _auto_width(ws2)

    # ── Databases ─────────────────────────────────────────────────────────────
    if result.databases:
        ws3 = wb.create_sheet("Databases")
        ws3.sheet_view.showGridLines = False
        headers = ["Name", "Owner", "Retention (days)", "Transient", "Created"]
        for ci, h in enumerate(headers, 1):
            _hdr(ws3, 1, ci, h)
        for i, d in enumerate(result.databases, start=2):
            shade = i % 2 == 0
            for ci, v in enumerate([
                d.name, d.owner or "", d.retention_time,
                "Yes" if d.is_transient else "No", d.created_on or "",
            ], 1):
                _cell(ws3, i, ci, v, shade)
        _auto_width(ws3)

    # ── Schemas ───────────────────────────────────────────────────────────────
    if result.schemas:
        ws4 = wb.create_sheet("Schemas")
        ws4.sheet_view.showGridLines = False
        headers = ["Database", "Schema", "Owner", "Managed Access", "Transient", "Retention"]
        for ci, h in enumerate(headers, 1):
            _hdr(ws4, 1, ci, h)
        for i, s in enumerate(result.schemas, start=2):
            shade = i % 2 == 0
            for ci, v in enumerate([
                s.database_name, s.name, s.owner or "",
                "Yes" if s.is_managed_access else "No",
                "Yes" if s.is_transient else "No",
                s.retention_time,
            ], 1):
                _cell(ws4, i, ci, v, shade)
        ws4.auto_filter.ref = f"A1:F{len(result.schemas)+1}"
        ws4.freeze_panes = "A2"
        _auto_width(ws4)

    # ── Tables ────────────────────────────────────────────────────────────────
    if result.tables:
        ws5 = wb.create_sheet("Tables & Views")
        ws5.sheet_view.showGridLines = False
        headers = ["Database", "Schema", "Name", "Type", "Rows", "Size (bytes)",
                   "Clustering Key", "Transient", "Created"]
        for ci, h in enumerate(headers, 1):
            _hdr(ws5, 1, ci, h)
        for i, t in enumerate(result.tables, start=2):
            shade = i % 2 == 0
            for ci, v in enumerate([
                t.database_name, t.schema_name, t.name, t.table_type,
                t.row_count or 0, t.bytes or 0, t.clustering_key or "",
                "Yes" if t.is_transient else "No", t.created or "",
            ], 1):
                _cell(ws5, i, ci, v, shade)
        ws5.auto_filter.ref = f"A1:I{len(result.tables)+1}"
        ws5.freeze_panes = "A2"
        _auto_width(ws5)

    # ── Object Inventory ──────────────────────────────────────────────────────
    if result.object_inventory:
        oi = result.object_inventory
        ws6 = wb.create_sheet("Object Inventory")
        ws6.sheet_view.showGridLines = False
        _hdr(ws6, 1, 1, "Object Type")
        _hdr(ws6, 1, 2, "Count")
        inventory_data = [
            ("Stages", oi.stages), ("Pipes (Snowpipe)", oi.pipes),
            ("Tasks (Scheduled)", oi.tasks), ("Streams (CDC)", oi.streams),
            ("Stored Procedures", oi.procedures), ("UDFs / Functions", oi.functions),
            ("Sequences", oi.sequences), ("File Formats", oi.file_formats),
            ("Dynamic Tables", oi.dynamic_tables),
            ("Data Shares (Outbound)", oi.shares_outbound),
            ("Data Shares (Inbound)", oi.shares_inbound),
            ("Resource Monitors", oi.resource_monitors),
            ("Network Policies", oi.network_policies),
            ("Masking Policies", oi.masking_policies),
            ("Row Access Policies", oi.row_access_policies),
        ]
        for i, (k, v) in enumerate(inventory_data, start=2):
            _cell(ws6, i, 1, k, i % 2 == 0)
            _cell(ws6, i, 2, v, i % 2 == 0)
        _auto_width(ws6)

    # ── Users ─────────────────────────────────────────────────────────────────
    if result.users:
        ws7 = wb.create_sheet("Users")
        ws7.sheet_view.showGridLines = False
        headers = ["Name", "Login Name", "Email", "Default Role", "Default WH",
                   "Disabled", "Has MFA", "Owner", "Last Login"]
        for ci, h in enumerate(headers, 1):
            _hdr(ws7, 1, ci, h)
        from openpyxl.styles import Font as OFont
        for i, u in enumerate(result.users, start=2):
            shade = i % 2 == 0
            row_vals = [
                u.get("name", ""), u.get("login_name", ""), u.get("email", ""),
                u.get("default_role", ""), u.get("default_warehouse", ""),
                "Yes" if u.get("disabled") else "No",
                "Yes" if u.get("has_mfa") else "No",
                u.get("owner", ""), u.get("last_success_login", ""),
            ]
            for ci, v in enumerate(row_vals, 1):
                c = ws7.cell(row=i, column=ci, value=v)
                c.font = _font()
                c.fill = _fill(SNOW_PALE if shade else LIGHT_GRAY)
                c.alignment = _align()
                c.border = _border()
                # Highlight users without MFA
                if ci == 7 and v == "No":
                    c.font = OFont(name=FONT_NAME, bold=True, color="EF4444")
        ws7.auto_filter.ref = f"A1:I{len(result.users)+1}"
        ws7.freeze_panes = "A2"
        _auto_width(ws7)

    # ── Roles ─────────────────────────────────────────────────────────────────
    if result.roles:
        ws8 = wb.create_sheet("Roles")
        ws8.sheet_view.showGridLines = False
        headers = ["Role Name", "Owner", "Assigned to Users", "Granted to Roles", "Comment"]
        for ci, h in enumerate(headers, 1):
            _hdr(ws8, 1, ci, h)
        for i, r in enumerate(result.roles, start=2):
            shade = i % 2 == 0
            for ci, v in enumerate([
                r.get("name", ""), r.get("owner", ""),
                r.get("assigned_to_users", 0), r.get("granted_to_roles", 0),
                r.get("comment", ""),
            ], 1):
                _cell(ws8, i, ci, v, shade)
        _auto_width(ws8)

    # ── Query Metrics ─────────────────────────────────────────────────────────
    if result.query_metrics and result.query_metrics.total_queries_last_7d > 0:
        qm = result.query_metrics
        ws9 = wb.create_sheet("Query Metrics (7d)")
        ws9.sheet_view.showGridLines = False
        _hdr(ws9, 1, 1, "Metric")
        _hdr(ws9, 1, 2, "Value")
        qm_data = [
            ("Total Queries (7 days)", qm.total_queries_last_7d),
            ("Failed Queries (7 days)", qm.failed_queries_last_7d),
            ("Avg Execution Time (ms)", round(qm.avg_execution_ms, 1)),
            ("P95 Execution Time (ms)", round(qm.p95_execution_ms, 1)),
            ("Total Bytes Scanned", qm.bytes_scanned_total),
            ("Bytes Spilled to Local Storage", qm.bytes_spilled_local),
            ("Bytes Spilled to Remote Storage", qm.bytes_spilled_remote),
        ]
        for i, (k, v) in enumerate(qm_data, start=2):
            _cell(ws9, i, 1, k, i % 2 == 0)
            _cell(ws9, i, 2, v, i % 2 == 0)
        _auto_width(ws9)

    # ── Cost / Credits ────────────────────────────────────────────────────────
    if result.cost_metrics and result.cost_metrics.credits_used_last_30d > 0:
        cm = result.cost_metrics
        ws10 = wb.create_sheet("Credits (30d)")
        ws10.sheet_view.showGridLines = False
        _hdr(ws10, 1, 1, "Metric")
        _hdr(ws10, 1, 2, "Value")
        cm_data = [
            ("Total Credits Used (30 days)", round(cm.credits_used_last_30d, 3)),
            ("Compute Credits", round(cm.compute_credits, 3)),
            ("Cloud Services Credits", round(cm.cloud_services_credits, 3)),
        ]
        for i, (k, v) in enumerate(cm_data, start=2):
            _cell(ws10, i, 1, k, i % 2 == 0)
            _cell(ws10, i, 2, v, i % 2 == 0)
        if cm.top_warehouses_by_credit:
            ws10.cell(row=len(cm_data) + 3, column=1, value="Top Warehouses by Credit Usage")
            _hdr(ws10, len(cm_data) + 4, 1, "Warehouse Name")
            _hdr(ws10, len(cm_data) + 4, 2, "Credits Used")
            for j, wh in enumerate(cm.top_warehouses_by_credit, start=len(cm_data) + 5):
                _cell(ws10, j, 1, wh.get("name", ""), j % 2 == 0)
                _cell(ws10, j, 2, round(wh.get("credits", 0), 3), j % 2 == 0)
        _auto_width(ws10)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Word report ───────────────────────────────────────────────────────────────

def _build_word(result: SnowflakeAssessmentResult) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    CYAN      = RGBColor(0x00, 0xB8, 0xE6)
    NAVY      = RGBColor(0x0A, 0x16, 0x28)
    WHITE_CLR = RGBColor(0xFF, 0xFF, 0xFF)
    GRAY      = RGBColor(0x44, 0x44, 0x44)

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)

    def _set_bg(cell, hex_c: str):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_c)
        tc_pr.append(shd)

    def _h(text: str, level: int = 1):
        p = doc.add_heading(text, level=level)
        p.runs[0].font.color.rgb = NAVY if level <= 2 else CYAN
        p.runs[0].font.bold = True

    def _kv_table(rows: list[tuple]):
        table = doc.add_table(rows=len(rows) + 1, cols=2)
        table.style = "Table Grid"
        hdr = table.rows[0]
        for i, h in enumerate(["Metric", "Value"]):
            hdr.cells[i].text = h
            hdr.cells[i].paragraphs[0].runs[0].font.bold = True
            hdr.cells[i].paragraphs[0].runs[0].font.color.rgb = WHITE_CLR
            hdr.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_bg(hdr.cells[i], SNOW_NAVY)
        for idx, (k, v) in enumerate(rows):
            row = table.rows[idx + 1]
            row.cells[0].text = str(k)
            row.cells[1].text = str(v)
            bg = SNOW_PALE.replace("#", "") if idx % 2 == 0 else "F5F5F5"
            _set_bg(row.cells[0], bg)
            _set_bg(row.cells[1], bg)
        doc.add_paragraph()

    # Cover
    acct = result.account_info.account_name if result.account_info else (result.label or "Unknown")
    doc.add_paragraph()
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run("SNOWFLAKE")
    r.font.size = Pt(32)
    r.font.bold = True
    r.font.color.rgb = CYAN

    sp = doc.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sp.add_run("Source Assessment Report")
    sr.font.size = Pt(18)
    sr.font.color.rgb = NAVY

    doc.add_paragraph()
    ap = doc.add_paragraph()
    ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ar = ap.add_run(acct)
    ar.font.size = Pt(14)
    ar.font.bold = True
    ar.font.color.rgb = GRAY

    dp = doc.add_paragraph()
    dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dp.add_run(f"Assessment Date: {result.assessed_at[:10]}").font.size = Pt(11)

    doc.add_page_break()

    # Executive Summary
    _h("Executive Summary", 1)
    ai = result.account_info
    ds = result.database_summary
    wm = result.warehouse_metrics
    up = result.user_profile
    if ai and ds and wm and up:
        doc.add_paragraph(
            f"Snowflake account '{ai.account_name}' (region: {ai.region or 'N/A'}) "
            f"is running version {ai.snowflake_version or 'unknown'}. "
            f"The account has {wm.total_warehouses} compute warehouse(s), "
            f"{ds.total_databases} databases, {ds.total_schemas} schemas, "
            f"{ds.total_tables} base tables, and {ds.total_views} views. "
            f"User base consists of {up.total_users} users across {up.total_roles} roles "
            f"({up.users_without_mfa} user(s) have no MFA configured)."
        )
    doc.add_paragraph()

    if ai:
        _h("Account Information", 2)
        _kv_table([
            ("Account Name", ai.account_name or ""),
            ("Account Locator", ai.account_locator or ""),
            ("Organization", ai.organization_name or ""),
            ("Region", ai.region or ""),
            ("Snowflake Version", ai.snowflake_version or ""),
            ("Current Role", ai.current_role or ""),
            ("Authenticated User", ai.current_user or ""),
        ])

    if wm:
        _h("Compute Warehouses", 2)
        _kv_table([
            ("Total Warehouses", wm.total_warehouses),
            ("Active (Started)", wm.active_warehouses),
            ("Suspended", wm.suspended_warehouses),
            ("Multi-Cluster Warehouses", wm.multi_cluster_warehouses),
        ])

    if ds:
        _h("Database & Object Summary", 2)
        _kv_table([
            ("Total Databases", ds.total_databases),
            ("Total Schemas", ds.total_schemas),
            ("Total Base Tables", ds.total_tables),
            ("Total Views", ds.total_views),
            ("External Tables", ds.total_external_tables),
            ("Materialized Views", ds.total_materialized_views),
            ("Total Data Size (bytes)", ds.total_size_bytes),
        ])

    if result.object_inventory:
        oi = result.object_inventory
        _h("Platform Object Inventory", 2)
        _kv_table([
            ("Stages", oi.stages),
            ("Snowpipe Pipes", oi.pipes),
            ("Scheduled Tasks", oi.tasks),
            ("CDC Streams", oi.streams),
            ("Stored Procedures", oi.procedures),
            ("UDFs / Functions", oi.functions),
            ("Dynamic Tables", oi.dynamic_tables),
            ("Data Shares (Outbound / Inbound)", f"{oi.shares_outbound} / {oi.shares_inbound}"),
            ("Masking Policies", oi.masking_policies),
            ("Row Access Policies", oi.row_access_policies),
        ])

    if up:
        _h("User & Role Profile", 2)
        _kv_table([
            ("Total Users", up.total_users),
            ("Disabled Users", up.disabled_users),
            ("Users Without MFA", up.users_without_mfa),
            ("Admin Users", up.admin_users),
            ("Service Accounts", up.service_accounts),
            ("Total Roles", up.total_roles),
            ("Custom Roles", up.custom_roles),
            ("System Roles", up.system_roles),
        ])

    if result.security_posture:
        sp2 = result.security_posture
        _h("Security Posture", 2)
        _kv_table([
            ("Network Policies", sp2.network_policies_count),
            ("Users Without MFA", sp2.users_without_mfa),
            ("Users with PUBLIC as Default Role", sp2.users_with_default_role_public),
            ("Masking Policies", sp2.masking_policies_count),
            ("Row Access Policies", sp2.row_access_policies_count),
            ("Data Shares (Total)", sp2.shares_total),
            ("Resource Monitors", sp2.resource_monitors_count),
        ])

    if result.query_metrics and result.query_metrics.total_queries_last_7d > 0:
        qm = result.query_metrics
        _h("Query Performance (Last 7 Days)", 2)
        _kv_table([
            ("Total Queries", qm.total_queries_last_7d),
            ("Failed Queries", qm.failed_queries_last_7d),
            ("Avg Execution Time (ms)", round(qm.avg_execution_ms, 1)),
            ("P95 Execution Time (ms)", round(qm.p95_execution_ms, 1)),
            ("Total Bytes Scanned", qm.bytes_scanned_total),
            ("Bytes Spilled to Local", qm.bytes_spilled_local),
            ("Bytes Spilled to Remote", qm.bytes_spilled_remote),
        ])

    if result.cost_metrics and result.cost_metrics.credits_used_last_30d > 0:
        cm = result.cost_metrics
        _h("Credit Usage (Last 30 Days)", 2)
        _kv_table([
            ("Total Credits Used", round(cm.credits_used_last_30d, 3)),
            ("Compute Credits", round(cm.compute_credits, 3)),
            ("Cloud Services Credits", round(cm.cloud_services_credits, 3)),
        ])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
