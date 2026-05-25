"""
AI-powered Word report builder — Microsoft Fabric Assessment Report (.docx)
Uses Azure OpenAI GPT-4o to generate narrative sections (Executive Summary,
Current State, Pain Points, Fabric Recommendations, ROI) and combines them
with data-driven tables populated directly from the assessment results.

Matches the style of Argon Medical Devices and CREA assessment reports.
"""

from __future__ import annotations

import io
import json
from datetime import datetime
from typing import Any

from docx import Document
from docx.shared import Cm, Pt
import openai

from app.config import settings
from app.core.logging import get_logger
from app.services.word_report_service import (
    _ACCENT_BLUE,
    _DARK_BLUE,
    _DARK_GRAY,
    _LIGHT_GRAY,
    _MID_BLUE,
    _TBL_HEADER,
    _TBL_ROW1,
    _TBL_ROW2,
    _WHITE,
    _body,
    _borders,
    _cell_vert_center,
    _cover_page,
    _font,
    _h1,
    _h2,
    _overview,
    _para_shd,
    _s,
    _setup_footer,
    _setup_header,
    _shd,
    _table,
)

logger = get_logger(__name__)

# ── System prompt — style guide from Argon + CREA sample reports ─────────────

_SYSTEM_PROMPT = """You are a senior Microsoft Fabric consultant at Unlimited Innovations India Pvt. Ltd.
You write professional database migration assessment reports recommending Microsoft Fabric adoption.

Writing style rules:
- Formal and executive-level tone
- Always cite specific numbers from the data provided (table counts, DB sizes, job counts, etc.)
- Concise paragraphs — no filler phrases like "It is important to note that..."
- Use Fabric/Microsoft terminology: Medallion Architecture, Bronze/Silver/Gold, OneLake, Lakehouse, Direct Lake, Data Factory, Fabric Pipelines, Microsoft Purview
- Forward-looking — highlight business value and ROI
- Never use first person ("we" is acceptable for the consulting firm, not "I")

Style examples from real Argon Medical Devices report:
"Argon Medical Devices operates a complex data ecosystem built on a QAD ERP system integrated with
Salesforce CRM, a bespoke BI database (ArgonBI), and an SSAS-based analytical layer serving 100–150
report consumers. The existing architecture relies heavily on tightly coupled SQL Agent jobs, SSIS packages,
and on-premises infrastructure, creating scalability bottlenecks and limiting real-time analytics capabilities."

Style examples from real CREA Financial report:
"CREA's current data landscape is anchored by two core SQL Server databases — PISCES (transactional,
~60 GB) and TAURUS (data warehouse, ~40 GB) — supplemented by Dataverse, SharePoint, and a legacy
MRI linked server. The platform supports approximately 125 users across finance, operations, and
executive reporting functions, with ~80 Power Automate flows managing critical business processes."
"""


# ── Azure OpenAI client ───────────────────────────────────────────────────────

def _get_client() -> openai.OpenAI:
    # Azure AI Foundry exposes an OpenAI-compatible endpoint at /openai/v1
    # Use the standard OpenAI client with the full base_url (not AzureOpenAI)
    return openai.OpenAI(
        base_url=settings.azure_openai_endpoint.rstrip("/"),
        api_key=settings.azure_openai_api_key.get_secret_value(),
    )


def _call_ai(client: openai.OpenAI, prompt: str, max_tokens: int = 800,
             json_mode: bool = False) -> str:
    kwargs: dict = dict(
        model=settings.azure_openai_deployment,
        messages=[
            {"role": "system", "content": _SYSTEM_PROMPT},
            {"role": "user", "content": prompt},
        ],
        max_tokens=max_tokens,
        temperature=0.3,
    )
    if json_mode:
        kwargs["response_format"] = {"type": "json_object"}
    response = client.chat.completions.create(**kwargs)
    return response.choices[0].message.content.strip()


# ── Context builder — extracts key metrics from raw assessment data ───────────

def _build_context(raw: dict, client_name: str) -> dict:
    ov = _overview(raw)

    pii        = raw.get("pii_indicators") or []
    orphaned   = raw.get("orphaned_users") or []
    dyn_sql    = raw.get("dynamic_sql_usage") or []
    linked     = raw.get("linked_servers") or []
    agent_jobs = raw.get("sql_agent_jobs") or []
    missing_ix = raw.get("missing_indexes") or []
    frag       = raw.get("fragmentation_report") or []
    tde        = raw.get("tde_status") or []
    clr        = raw.get("clr_assemblies") or []
    cross_db   = raw.get("cross_db_references") or []
    deprecated = raw.get("deprecated_features_in_use") or []
    heaps      = raw.get("heap_tables") or []
    missing_pk = raw.get("missing_primary_keys") or []
    tables     = raw.get("tables") or []
    vf         = raw.get("version_features") or []
    ssis       = (raw.get("ssis_catalog_packages") or []) + (raw.get("ssis_msdb_packages") or [])
    users_roles = raw.get("db_users_roles") or []
    null_a     = raw.get("null_analysis") or []

    size_mb = ov.get("total_size_mb")
    size_gb = f"{float(size_mb) / 1024:.2f}" if size_mb else "Unknown"

    version_str = _s(vf[0].get("product_version") if vf else None, "Unknown")
    edition     = _s(vf[0].get("edition") if vf else None, "Unknown")

    tde_enabled = any(
        str(t.get("encryption_state") or t.get("tde_status", "")).lower()
        in ("3", "encrypted", "enabled")
        for t in tde
    )

    top_tables = sorted(
        [t for t in tables if t.get("row_count")],
        key=lambda x: int(x.get("row_count") or 0),
        reverse=True,
    )[:10]

    return {
        "client_name":             client_name,
        "database_name":           _s(ov.get("database_name"), "Unknown"),
        "server_version":          version_str,
        "edition":                 edition,
        "database_size_gb":        size_gb,
        "table_count":             _s(ov.get("table_count"), "0"),
        "view_count":              _s(ov.get("view_count"), "0"),
        "stored_proc_count":       _s(ov.get("stored_proc_count"), "0"),
        "function_count":          _s(ov.get("function_count"), "0"),
        "schema_count":            _s(ov.get("schema_count"), "0"),
        "pii_column_count":        len(pii),
        "pii_samples":             [
            f"{p.get('schema_name','')}.{p.get('table_name','')}.{p.get('column_name','')} ({p.get('pii_category','PII')})"
            for p in pii[:5]
        ],
        "orphaned_user_count":     len(orphaned),
        "dynamic_sql_count":       len(dyn_sql),
        "linked_server_count":     len(linked),
        "linked_server_names":     [_s(ls.get("linked_server_name")) for ls in linked[:5]],
        "agent_job_count":         len(agent_jobs),
        "agent_job_names":         [_s(j.get("job_name")) for j in agent_jobs[:10]],
        "missing_index_count":     len(missing_ix),
        "fragmented_table_count":  len(frag),
        "tde_enabled":             tde_enabled,
        "clr_assembly_count":      len(clr),
        "cross_db_count":          len(cross_db),
        "deprecated_feature_count": len(deprecated),
        "heap_table_count":        len(heaps),
        "missing_pk_count":        len(missing_pk),
        "ssis_package_count":      len(ssis),
        "user_role_count":         len(users_roles),
        "null_analysis_count":     len(null_a),
        "top_tables":              top_tables,
        # raw slices for table rendering
        "_pii":         pii,
        "_orphaned":    orphaned,
        "_linked":      linked,
        "_agent_jobs":  agent_jobs,
        "_missing_ix":  missing_ix,
        "_frag":        frag,
        "_tde":         tde,
        "_users_roles": users_roles,
        "_null_a":      null_a,
        "_dyn_sql":     dyn_sql,
    }


# ── AI generation functions ───────────────────────────────────────────────────

def _gen_executive_summary(client: openai.OpenAI, ctx: dict) -> str:
    prompt = f"""Write a 2-paragraph Executive Summary for a Microsoft Fabric Assessment Report.

Assessment data:
- Client: {ctx['client_name']}
- Database: {ctx['database_name']} | Size: {ctx['database_size_gb']} GB
- SQL Server: {ctx['server_version']} ({ctx['edition']})
- Schemas: {ctx['schema_count']} | Tables: {ctx['table_count']} | Views: {ctx['view_count']}
- Stored Procedures: {ctx['stored_proc_count']} | Functions: {ctx['function_count']}
- SQL Agent Jobs: {ctx['agent_job_count']} | SSIS Packages: {ctx['ssis_package_count']}
- Linked Servers: {ctx['linked_server_count']} ({', '.join(ctx['linked_server_names']) or 'none'})
- PII columns detected: {ctx['pii_column_count']}
- Missing indexes: {ctx['missing_index_count']} | Fragmented tables: {ctx['fragmented_table_count']}
- Heap tables (no clustered index): {ctx['heap_table_count']}
- Orphaned users: {ctx['orphaned_user_count']} | Dynamic SQL objects: {ctx['dynamic_sql_count']}
- TDE enabled: {ctx['tde_enabled']} | CLR assemblies: {ctx['clr_assembly_count']}
- Cross-database references: {ctx['cross_db_count']} | Deprecated features: {ctx['deprecated_feature_count']}

Paragraph 1: Describe the current data landscape — systems involved, scale, complexity, user base.
Paragraph 2: Summarize the key challenges identified and the Microsoft Fabric migration recommendation.

Rules: cite specific numbers; formal consultant tone; no bullet points; 4–6 sentences per paragraph."""
    return _call_ai(client, prompt, max_tokens=700)


def _gen_current_state(client: openai.OpenAI, ctx: dict) -> str:
    prompt = f"""Write a Current State Overview paragraph (4–5 sentences) for {ctx['client_name']}'s data platform.

Facts:
- Core database: {ctx['database_name']} ({ctx['database_size_gb']} GB, {ctx['table_count']} tables, {ctx['view_count']} views)
- ETL layer: {ctx['agent_job_count']} SQL Agent Jobs, {ctx['ssis_package_count']} SSIS packages
- Linked systems: {ctx['linked_server_count']} linked server(s) ({', '.join(ctx['linked_server_names']) or 'none detected'})
- Reporting: views and stored procedures ({ctx['view_count']} views, {ctx['stored_proc_count']} procs)
- Users/roles configured: {ctx['user_role_count']}

Describe the architecture as-is: data sources, processing layer, and reporting/consumption layer.
Formal consultant tone. Cite the numbers. No bullet points."""
    return _call_ai(client, prompt, max_tokens=400)


def _gen_pain_points(client: openai.OpenAI, ctx: dict) -> list[dict]:
    prompt = f"""Identify 8–10 key pain points for a Microsoft Fabric migration assessment report.

Assessment findings for {ctx['client_name']} — {ctx['database_name']} ({ctx['database_size_gb']} GB):
- SQL Agent Jobs: {ctx['agent_job_count']} ({', '.join(ctx['agent_job_names'][:5]) or 'none'})
- SSIS Packages: {ctx['ssis_package_count']}
- PII columns: {ctx['pii_column_count']} ({', '.join(ctx['pii_samples'][:3]) or 'none detected'})
- Orphaned users: {ctx['orphaned_user_count']}
- Dynamic SQL objects: {ctx['dynamic_sql_count']}
- Missing indexes: {ctx['missing_index_count']}
- Fragmented tables: {ctx['fragmented_table_count']}
- Linked servers: {ctx['linked_server_count']} ({', '.join(ctx['linked_server_names']) or 'none'})
- Heap tables (no clustered index): {ctx['heap_table_count']}
- Missing primary keys: {ctx['missing_pk_count']}
- Deprecated features in use: {ctx['deprecated_feature_count']}
- CLR assemblies: {ctx['clr_assembly_count']}
- Cross-DB references: {ctx['cross_db_count']}
- TDE enabled: {ctx['tde_enabled']}

Return a JSON object with key "pain_points" containing an array of objects.
Each object must have exactly these three string fields:
  "category" - short category (e.g. "ETL Complexity", "Data Security", "Performance", "Governance")
  "issue"    - specific issue found, 1–2 sentences, cite numbers from the data above
  "impact"   - business impact if not addressed, 1 sentence

Only return valid JSON. No markdown fences."""
    raw_text = _call_ai(client, prompt, max_tokens=1800, json_mode=True)
    try:
        parsed = json.loads(raw_text)
        items = parsed.get("pain_points") or parsed.get("items") or []
        if isinstance(items, list) and items:
            return items
    except (json.JSONDecodeError, ValueError, TypeError):
        pass

    # Fallback static pain points
    return [
        {"category": "Legacy ETL Infrastructure",
         "issue": f"{ctx['agent_job_count']} SQL Agent Jobs orchestrate data movement with no centralised monitoring, alerting, or lineage tracking.",
         "impact": "Undetected job failures cause stale reports and require manual intervention, increasing operational overhead."},
        {"category": "SSIS Dependency",
         "issue": f"{ctx['ssis_package_count']} SSIS packages introduce on-premises infrastructure dependency, limiting cloud scalability.",
         "impact": "Any infrastructure change risks pipeline breakage, slowing delivery of new data integrations."},
        {"category": "Performance Degradation",
         "issue": f"{ctx['missing_index_count']} missing indexes and {ctx['fragmented_table_count']} fragmented tables have been identified, degrading query performance.",
         "impact": "Report execution times increase, reducing analyst productivity and delaying business decisions."},
        {"category": "Data Security & PII",
         "issue": f"{ctx['pii_column_count']} PII-bearing columns lack consistent data masking or sensitivity classification.",
         "impact": "Regulatory compliance risk under GDPR and local data privacy laws if sensitive data is exposed."},
        {"category": "Governance Gaps",
         "issue": "No centralised data catalog or end-to-end lineage exists; data consumers cannot trace report figures back to source.",
         "impact": "Erodes trust in data assets and complicates audit and compliance reporting."},
        {"category": "Scalability Constraints",
         "issue": f"The {ctx['database_size_gb']} GB database runs on fixed on-premises infrastructure with no elastic scaling capability.",
         "impact": "Peak-load periods cause resource contention and report timeouts, impacting user experience."},
        {"category": "Security Hygiene",
         "issue": f"{ctx['orphaned_user_count']} orphaned database users and {ctx['dynamic_sql_count']} objects using dynamic SQL were detected.",
         "impact": "Orphaned accounts expand the attack surface; dynamic SQL increases SQL injection risk."},
        {"category": "Modernisation Readiness",
         "issue": f"{ctx['deprecated_feature_count']} deprecated SQL Server features in active use will not be supported in future versions.",
         "impact": "Platform upgrades are blocked until deprecated features are refactored, increasing technical debt."},
    ]


def _gen_fabric_recommendations(client: openai.OpenAI, ctx: dict) -> str:
    prompt = f"""Write a 3-paragraph Fabric Recommendations section for {ctx['client_name']}'s migration assessment.

Current state:
- Database: {ctx['database_name']}, {ctx['database_size_gb']} GB, {ctx['table_count']} tables
- ETL: {ctx['agent_job_count']} SQL Agent Jobs, {ctx['ssis_package_count']} SSIS packages
- Linked servers: {ctx['linked_server_count']} | Cross-DB references: {ctx['cross_db_count']}
- Views: {ctx['view_count']} | Stored procs: {ctx['stored_proc_count']}
- PII columns: {ctx['pii_column_count']} | Missing indexes: {ctx['missing_index_count']}

Paragraph 1: Proposed Microsoft Fabric target architecture.
  - Medallion Architecture (Bronze = raw ingestion via Data Factory, Silver = cleansed/conformed, Gold = semantic models)
  - Replace SQL Agent Jobs / SSIS with Fabric Pipelines and Dataflows Gen2
  - Lakehouse as the central storage layer (OneLake)

Paragraph 2: Specific migration priorities for this client (reference their actual numbers/systems).
  - Migrate the {ctx['database_name']} database to Fabric Lakehouse
  - Rebuild {ctx['agent_job_count']} Agent Jobs as Fabric Pipelines
  - Migrate views/reports to Direct Lake semantic models
  - Address {ctx['pii_column_count']} PII columns with Purview sensitivity labels

Paragraph 3: Expected outcomes — performance, cost, governance, and AI enablement.

Formal consultant tone. Be specific to this client's data. No bullet points in the paragraphs."""
    return _call_ai(client, prompt, max_tokens=800)


def _gen_roi_narrative(client: openai.OpenAI, ctx: dict) -> str:
    prompt = f"""Write a short ROI & Business Value paragraph (4–5 sentences) for migrating
{ctx['client_name']}'s {ctx['database_name']} database ({ctx['database_size_gb']} GB, {ctx['table_count']} tables)
to Microsoft Fabric.

Include:
- Estimated infrastructure cost reduction (typically 25–35% for SQL Server on-premises → Fabric)
- Reduction in ETL maintenance (replacing {ctx['agent_job_count']} SQL Agent Jobs and {ctx['ssis_package_count']} SSIS packages)
- Time-to-insight improvement via Direct Lake mode (no import refresh delays)
- Compliance and governance value from addressing {ctx['pii_column_count']} unclassified PII columns
- AI/ML enablement through Fabric notebooks and AutoML capabilities

Formal consultant tone. Cite the numbers. 1 paragraph only."""
    return _call_ai(client, prompt, max_tokens=350)


def _gen_conclusion(client: openai.OpenAI, ctx: dict) -> str:
    prompt = f"""Write a Conclusion paragraph (3–4 sentences) summarising the assessment findings
and reinforcing the Microsoft Fabric migration recommendation for {ctx['client_name']}.

Key facts: {ctx['database_name']} database ({ctx['database_size_gb']} GB), {ctx['table_count']} tables,
{ctx['agent_job_count']} SQL Agent Jobs, {ctx['pii_column_count']} PII columns,
{ctx['missing_index_count']} missing indexes, {ctx['fragmented_table_count']} fragmented tables.

Be decisive and forward-looking. Formal tone. No bullet points."""
    return _call_ai(client, prompt, max_tokens=300)


# ── Document section builders ─────────────────────────────────────────────────

def _write_executive_summary(doc: Document, text: str) -> None:
    _h1(doc, "Executive Summary")
    for para in text.split("\n\n"):
        if para.strip():
            _body(doc, para.strip())


def _write_current_state(doc: Document, text: str, ctx: dict) -> None:
    _h1(doc, "Current State Overview")
    _body(doc, text)

    _h2(doc, "Database Overview")
    db_type = (
        "Datawarehouse"
        if any(k in ctx["database_name"].lower() for k in ("dw", "warehouse", "mart", "dwh"))
        else "Transaction"
    )
    _table(
        doc,
        ["Database Name", "Size (GB)", "Type", "Tables", "Views", "Stored Procs", "Functions"],
        [[
            ctx["database_name"],
            ctx["database_size_gb"],
            db_type,
            ctx["table_count"],
            ctx["view_count"],
            ctx["stored_proc_count"],
            ctx["function_count"],
        ]],
        [3.5, 2, 2.5, 2, 2, 2.5, 2.5],
    )

    # Top tables by row volume
    if ctx["top_tables"]:
        _h2(doc, "High-Volume Tables")
        top_rows = [
            [
                f"{t.get('schema_name','dbo')}.{t.get('table_name','')}",
                f"{int(t.get('row_count') or 0):,}",
                _s(t.get("data_space_mb") or t.get("reserved_mb"), "—"),
            ]
            for t in ctx["top_tables"]
        ]
        _table(doc, ["Table Name", "Row Count", "Size (MB)"], top_rows, [8, 4, 4])

    # Linked servers
    if ctx["_linked"]:
        _h2(doc, "Linked Servers / External Connections")
        ls_rows = [
            [_s(r.get("linked_server_name")), _s(r.get("product") or r.get("provider")), _s(r.get("data_source"))]
            for r in ctx["_linked"]
        ]
        _table(doc, ["Linked Server", "Product / Provider", "Data Source"], ls_rows, [5, 5, 6])

    # SQL Agent Jobs
    if ctx["_agent_jobs"]:
        _h2(doc, "SQL Agent Jobs")
        job_rows = [
            [_s(j.get("job_name")), _s(j.get("enabled") or "—"), _s(j.get("description") or "—")]
            for j in ctx["_agent_jobs"][:20]
        ]
        _table(doc, ["Job Name", "Enabled", "Description"], job_rows, [7, 2.5, 6.5])


def _write_pain_points(doc: Document, pain_points: list[dict]) -> None:
    _h1(doc, "Key Pain Points")
    rows = [
        [p.get("category", "—"), p.get("issue", "—"), p.get("impact", "—")]
        for p in pain_points
    ]
    _table(doc, ["Category", "Key Issue", "Business Impact"], rows, [4, 7, 5])


def _write_security_section(doc: Document, ctx: dict) -> None:
    _h1(doc, "Security & Compliance Assessment")

    _h2(doc, "PII Indicators")
    if ctx["_pii"]:
        pii_rows = [
            [
                _s(r.get("pii_category") or "PII"),
                _s(r.get("schema_name")),
                _s(r.get("table_name")),
                _s(r.get("column_name")),
            ]
            for r in ctx["_pii"][:30]
        ]
        _table(
            doc,
            ["PII Category", "Schema", "Table", "Column"],
            pii_rows,
            [3.5, 3, 4, 5.5],
        )
    else:
        _body(doc, "No PII indicators detected in this assessment.")

    _h2(doc, "Database Users & Roles")
    if ctx["_users_roles"]:
        ur_rows = [
            [
                _s(r.get("principal_name") or r.get("user_name")),
                _s(r.get("principal_type") or "DB User"),
                _s(r.get("roles") or "—"),
            ]
            for r in ctx["_users_roles"][:20]
        ]
        _table(doc, ["User / Principal", "Type", "Roles"], ur_rows, [5, 3.5, 7.5])
    else:
        _body(doc, "No user/role data available at the current access level.")

    _h2(doc, "Orphaned Users")
    if ctx["_orphaned"]:
        orp_rows = [
            [_s(r.get("user_name") or r.get("principal_name")), _s(r.get("type_desc") or "—")]
            for r in ctx["_orphaned"]
        ]
        _table(doc, ["Orphaned User", "Type"], orp_rows, [8, 8])
    else:
        _body(doc, "No orphaned users detected.")

    _h2(doc, "Encryption Status (TDE)")
    if ctx["_tde"]:
        tde_rows = [
            [
                _s(r.get("database_name") or ctx["database_name"]),
                _s(r.get("encryption_state") or r.get("tde_status")),
                _s(r.get("key_algorithm") or "—"),
                _s(r.get("key_length") or "—"),
            ]
            for r in ctx["_tde"]
        ]
        _table(doc, ["Database", "Encryption State", "Algorithm", "Key Length"], tde_rows, [4, 4, 4, 4])
    else:
        _body(doc, "TDE status not available at the current access level.")

    _h2(doc, "Dynamic SQL Usage")
    if ctx["_dyn_sql"]:
        dyn_rows = [
            [
                _s(r.get("schema_name")),
                _s(r.get("object_name")),
                _s(r.get("object_type") or "—"),
            ]
            for r in ctx["_dyn_sql"][:20]
        ]
        _table(doc, ["Schema", "Object Name", "Type"], dyn_rows, [4, 8, 4])
    else:
        _body(doc, "No dynamic SQL usage detected.")


def _write_performance_section(doc: Document, ctx: dict) -> None:
    _h1(doc, "Performance Analysis")

    _h2(doc, "Missing Index Recommendations")
    if ctx["_missing_ix"]:
        ix_rows = [
            [
                _s(r.get("table_name") or r.get("object_name")),
                _s(r.get("equality_columns") or r.get("index_columns") or "—"),
                _s(r.get("avg_user_impact") or r.get("impact") or "—"),
                _s(r.get("unique_compiles") or "—"),
            ]
            for r in ctx["_missing_ix"][:20]
        ]
        _table(
            doc,
            ["Table", "Suggested Columns", "Avg Impact (%)", "Compiles"],
            ix_rows,
            [5, 6, 3, 2],
        )
    else:
        _body(doc, "No missing index recommendations at the current access level.")

    _h2(doc, "Index Fragmentation")
    if ctx["_frag"]:
        frag_rows = [
            [
                _s(r.get("table_name") or r.get("object_name")),
                _s(r.get("index_name")),
                _s(r.get("avg_fragmentation_in_percent") or r.get("fragmentation_pct") or "—"),
                _s(r.get("page_count") or "—"),
            ]
            for r in ctx["_frag"][:20]
        ]
        _table(
            doc,
            ["Table", "Index Name", "Fragmentation (%)", "Pages"],
            frag_rows,
            [5, 5, 3.5, 2.5],
        )
    else:
        _body(doc, "No fragmentation data available at the current access level.")

    _h2(doc, "Null Analysis Summary")
    if ctx["_null_a"]:
        null_rows = [
            [
                _s(r.get("table_name")),
                _s(r.get("column_name")),
                _s(r.get("null_blank_pct") or r.get("null_percentage") or "—"),
                _s(r.get("total_rows") or "—"),
            ]
            for r in ctx["_null_a"][:20]
        ]
        _table(
            doc,
            ["Table", "Column", "Null %", "Total Rows"],
            null_rows,
            [5, 5, 2.5, 3.5],
        )
    else:
        _body(doc, "Null analysis was not run or no data is available.")


def _write_fabric_recommendations(doc: Document, text: str, ctx: dict) -> None:
    _h1(doc, "Fabric Recommendations")

    for para in text.split("\n\n"):
        if para.strip():
            _body(doc, para.strip())

    doc.add_paragraph()

    _h2(doc, "Proposed Medallion Architecture")
    arch_rows = [
        ["Bronze (Raw)",  "OneLake / Lakehouse",
         f"Ingest {ctx['database_name']} tables via Data Factory pipelines (full + incremental)"],
        ["Silver (Curated)", "Lakehouse / Warehouse",
         f"Cleanse and conform {ctx['table_count']} tables; replace {ctx['stored_proc_count']} stored procedures with Dataflows Gen2"],
        ["Gold (Serving)", "Semantic Model (Direct Lake)",
         f"Rebuild {ctx['view_count']} views as Direct Lake semantic model; serve Power BI reports"],
    ]
    _table(
        doc,
        ["Layer", "Fabric Component", "Migration Action"],
        arch_rows,
        [2.5, 4, 9.5],
    )

    _h2(doc, "ETL Modernisation")
    etl_rows = [
        ["SQL Agent Jobs", str(ctx["agent_job_count"]), "Fabric Pipelines / Data Factory", "High"],
        ["SSIS Packages",  str(ctx["ssis_package_count"]), "Dataflows Gen2 / Fabric Pipelines", "High"],
        ["Linked Servers", str(ctx["linked_server_count"]), "Fabric Shortcuts / Data Gateway", "Medium"],
        ["Stored Procedures", ctx["stored_proc_count"], "Dataflows Gen2 / Notebooks", "Medium"],
    ]
    _table(
        doc,
        ["Current Component", "Count", "Fabric Replacement", "Priority"],
        etl_rows,
        [4, 2, 6, 4],
    )

    _h2(doc, "Data Science & AI Enablement")
    _table(
        doc,
        ["Scenario", "Description", "Fabric Capability"],
        [
            ["Predictive Analytics",   "Use historical transaction data for demand forecasting and anomaly detection", "Fabric Notebooks + AutoML"],
            ["Real-Time Intelligence", "Trigger alerts on data quality thresholds and operational KPI breaches",      "Data Activator (Reflex)"],
            ["Natural Language Query", "Enable business users to query data in plain English via Copilot",           "Microsoft Fabric Copilot"],
        ],
        [4, 8, 4],
    )

    _h2(doc, "Governance & Security")
    gov_rows = [
        ["Sensitivity Labels",   f"Classify {ctx['pii_column_count']} PII columns",          "Microsoft Purview", "Yes"],
        ["Data Lineage",         "Source → Lakehouse → Semantic Model → Report",              "Purview Lineage",   "Yes"],
        ["Row-Level Security",   "Enforce existing role permissions in Fabric semantic model", "Power BI RLS",      "Yes"],
        ["Orphaned Accounts",    f"Remediate {ctx['orphaned_user_count']} orphaned DB users", "Entra ID Review",   "Yes"],
    ]
    _table(
        doc,
        ["Activity", "Detail", "Tool", "Recommended"],
        gov_rows,
        [4, 6, 4, 2],
    )


def _write_roi_section(doc: Document, narrative: str, ctx: dict) -> None:
    _h1(doc, "ROI & Business Value")

    for para in narrative.split("\n\n"):
        if para.strip():
            _body(doc, para.strip())

    doc.add_paragraph()

    _h2(doc, "Cost & Effort Comparison")
    _table(
        doc,
        ["Category", "Current State", "Post-Fabric", "Expected Saving"],
        [
            ["Infrastructure",       "On-premises SQL Server licensing + hardware", "Fabric capacity SKU (F64 recommended)", "~30% reduction"],
            ["ETL Maintenance",      f"{ctx['agent_job_count']} Agent Jobs + {ctx['ssis_package_count']} SSIS packages (manual ops)", "Fabric Pipelines with built-in monitoring", "~40% effort reduction"],
            ["Report Refresh Latency", "Scheduled import refresh (30–60 min lag)", "Direct Lake (near real-time)", "Near-zero refresh lag"],
            ["Governance & Compliance", f"{ctx['pii_column_count']} unclassified PII columns — manual process", "Purview auto-classification + sensitivity labels", "Audit risk eliminated"],
        ],
        [3, 5, 5, 3],
    )

    _h2(doc, "Implementation Timeline")
    _table(
        doc,
        ["Phase", "Focus Area", "Duration", "Key Deliverables"],
        [
            ["Phase 1", "Foundation & Connectivity",     "2 weeks", "Fabric workspace, OneLake, gateway setup"],
            ["Phase 2", "Bronze Layer — Raw Ingestion",  "3 weeks", f"Ingest all {ctx['table_count']} tables to Lakehouse"],
            ["Phase 3", "Silver Layer — Transformation", "3 weeks", f"Rebuild {ctx['stored_proc_count']} procs as Dataflows Gen2"],
            ["Phase 4", "Gold Layer — Semantic Model",   "2 weeks", f"Direct Lake model from {ctx['view_count']} views"],
            ["Phase 5", "Reporting & Validation",        "2 weeks", "Power BI migration, data reconciliation"],
            ["Phase 6", "Cutover & Governance",          "2 weeks", "Purview labels, RLS, monitoring, handover"],
        ],
        [2, 4, 2.5, 7.5],
    )


def _write_conclusion(doc: Document, text: str, ctx: dict) -> None:
    _h1(doc, "Conclusion")

    _h2(doc, "Summary of Findings")
    summary_rows = [
        ["Database Name",          ctx["database_name"]],
        ["SQL Server Version",     ctx["server_version"]],
        ["Edition",                ctx["edition"]],
        ["Database Size",          f"{ctx['database_size_gb']} GB"],
        ["Total Tables",           ctx["table_count"]],
        ["Total Views",            ctx["view_count"]],
        ["Stored Procedures",      ctx["stored_proc_count"]],
        ["Functions",              ctx["function_count"]],
        ["SQL Agent Jobs",         str(ctx["agent_job_count"])],
        ["SSIS Packages",          str(ctx["ssis_package_count"])],
        ["PII Columns Detected",   str(ctx["pii_column_count"])],
        ["Missing Indexes",        str(ctx["missing_index_count"])],
        ["Fragmented Tables",      str(ctx["fragmented_table_count"])],
        ["Orphaned Users",         str(ctx["orphaned_user_count"])],
        ["Heap Tables",            str(ctx["heap_table_count"])],
        ["Missing Primary Keys",   str(ctx["missing_pk_count"])],
        ["TDE Enabled",            "Yes" if ctx["tde_enabled"] else "No"],
        ["CLR Assemblies",         str(ctx["clr_assembly_count"])],
        ["Deprecated Features",    str(ctx["deprecated_feature_count"])],
    ]
    _table(doc, ["Metric", "Value"], summary_rows, [8, 8])

    _h2(doc, "Final Recommendations")
    for para in text.split("\n\n"):
        if para.strip():
            _body(doc, para.strip())

    _h2(doc, "Appendices")
    _body(
        doc,
        "Full schema details, index coverage, null analysis, and raw query results "
        "are available in the companion Excel workbook generated by the Source Assessment Tool.",
    )


# ═══════════════════════════════════════════════════════════════════════════════
# PUBLIC API
# ═══════════════════════════════════════════════════════════════════════════════

def build_ai_word_report(
    job_id: str,
    raw: dict[str, Any],
    client_name: str | None = None,
) -> bytes:
    """
    Build an AI-powered Word assessment report for a single-database assessment.
    Uses Azure OpenAI GPT-4o to generate Executive Summary, Current State, Pain Points,
    Fabric Recommendations, ROI narrative, and Conclusion.
    Data tables (DB metrics, security, performance) are populated from assessment results.
    """
    ov          = _overview(raw)
    db_name     = _s(ov.get("database_name"), "Unknown Database")
    server_name = _s(raw.get("_server") or ov.get("server_name"), "")
    label       = client_name or db_name
    run_date    = datetime.utcnow().strftime("%d %b %Y")

    ctx = _build_context(raw, label)
    ai  = _get_client()

    logger.info("AI report: generating content for job %s (client: %s)", job_id, label)

    executive_summary     = _gen_executive_summary(ai, ctx)
    current_state         = _gen_current_state(ai, ctx)
    pain_points           = _gen_pain_points(ai, ctx)
    fabric_recommendations = _gen_fabric_recommendations(ai, ctx)
    roi_narrative         = _gen_roi_narrative(ai, ctx)
    conclusion            = _gen_conclusion(ai, ctx)

    logger.info("AI report: content ready for job %s — assembling Word document", job_id)

    doc = Document()
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    _setup_header(doc, label)
    _setup_footer(doc)
    _cover_page(doc, label, run_date)

    _write_executive_summary(doc, executive_summary)
    _write_current_state(doc, current_state, ctx)
    _write_pain_points(doc, pain_points)
    _write_security_section(doc, ctx)
    _write_performance_section(doc, ctx)
    _write_fabric_recommendations(doc, fabric_recommendations, ctx)
    _write_roi_section(doc, roi_narrative, ctx)
    _write_conclusion(doc, conclusion, ctx)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    logger.info("AI report: Word document built for job %s (%s)", job_id, label)
    return buf.getvalue()


def build_ai_session_word_report(
    session_id: str,
    session_label: str | None,
    jobs_data: list[dict],
) -> bytes:
    """
    Build a combined AI-powered Word report for a multi-database session.
    Generates one set of AI narrative per database, then combines into one document.
    jobs_data: list of {job_id, server, database, label, results: raw_dict}
    """
    client_name = session_label or f"Session {session_id[:8]}"
    run_date    = datetime.utcnow().strftime("%d %b %Y")
    ai          = _get_client()

    doc = Document()
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    _setup_header(doc, client_name)
    _setup_footer(doc)
    _cover_page(doc, client_name, run_date)

    # Session summary table
    _h1(doc, "Session Overview")
    _body(doc, f"This report covers {len(jobs_data)} database assessment(s) under session: {client_name}")
    doc.add_paragraph()

    sum_rows = []
    for jd in jobs_data:
        raw = jd.get("results") or {}
        ov  = _overview(raw)
        sz  = ov.get("total_size_mb")
        sum_rows.append([
            _s(jd.get("server")),
            _s(jd.get("database")),
            _s(ov.get("table_count")),
            _s(ov.get("view_count")),
            _s(ov.get("stored_proc_count")),
            f"{float(sz)/1024:.2f} GB" if sz else "—",
        ])
    _table(
        doc,
        ["Server", "Database", "Tables", "Views", "Stored Procs", "Size (GB)"],
        sum_rows,
        [4.5, 3.5, 2, 2, 2.5, 2.5],
    )

    for idx, jd in enumerate(jobs_data):
        raw    = jd.get("results") or {}
        db_lbl = _s(jd.get("database"), f"Database {idx + 1}")
        srv_lbl = _s(jd.get("server"))
        label  = jd.get("label") or db_lbl

        doc.add_page_break()

        banner = doc.add_paragraph()
        _para_shd(banner, _MID_BLUE)
        br = banner.add_run(f"  Database: {srv_lbl}  ›  {db_lbl}")
        _font(br, bold=True, size=14, colour=_WHITE)
        banner.paragraph_format.space_before = Pt(0)
        banner.paragraph_format.space_after  = Pt(10)

        ctx = _build_context(raw, label)

        logger.info("AI session report: generating content for %s / %s", srv_lbl, db_lbl)

        executive_summary      = _gen_executive_summary(ai, ctx)
        current_state          = _gen_current_state(ai, ctx)
        pain_points            = _gen_pain_points(ai, ctx)
        fabric_recommendations = _gen_fabric_recommendations(ai, ctx)
        roi_narrative          = _gen_roi_narrative(ai, ctx)
        conclusion             = _gen_conclusion(ai, ctx)

        _write_executive_summary(doc, executive_summary)
        _write_current_state(doc, current_state, ctx)
        _write_pain_points(doc, pain_points)
        _write_security_section(doc, ctx)
        _write_performance_section(doc, ctx)
        _write_fabric_recommendations(doc, fabric_recommendations, ctx)
        _write_roi_section(doc, roi_narrative, ctx)
        _write_conclusion(doc, conclusion, ctx)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    logger.info("AI session report: complete for session %s (%d databases)", session_id, len(jobs_data))
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════════
# FABRIC ASSESSMENT WORD REPORT
# ═══════════════════════════════════════════════════════════════════════════════

def _build_fabric_context(results: dict, client_name: str) -> dict:
    """Extract key metrics from Fabric session results into a concise context dict."""
    summary    = results.get("summary") or {}
    workspaces = results.get("workspaces") or []

    all_datasets = [ds for ws in workspaces for ds in (ws.get("datasets") or [])]
    all_reports  = [rp for ws in workspaces for rp in (ws.get("reports") or [])]
    all_measures = [m for ds in all_datasets for m in (ds.get("measures") or [])]
    all_tables   = [t for ds in all_datasets for t in (ds.get("tables") or [])]

    # Complexity distribution
    complexity_dist: dict[str, int] = {}
    for m in all_measures:
        level = (m.get("complexity") or {}).get("level") or "Unknown"
        complexity_dist[level] = complexity_dist.get(level, 0) + 1

    top_complex = sorted(
        [m for m in all_measures if (m.get("complexity") or {}).get("score", 0) > 0],
        key=lambda x: (x.get("complexity") or {}).get("score", 0),
        reverse=True,
    )[:10]

    storage_modes: dict[str, int] = {}
    for ds in all_datasets:
        mode = ds.get("storage_mode") or "Unknown"
        storage_modes[mode] = storage_modes.get(mode, 0) + 1

    return {
        "client_name":        client_name,
        "workspace_count":    summary.get("workspace_count", len(workspaces)),
        "dataset_count":      summary.get("dataset_count", len(all_datasets)),
        "report_count":       summary.get("report_count", len(all_reports)),
        "total_measures":     summary.get("total_measures", len(all_measures)),
        "total_tables":       len(all_tables),
        "total_visuals":      summary.get("total_visuals", 0),
        "total_relationships": summary.get("total_relationships", 0),
        "complexity_dist":    complexity_dist,
        "storage_modes":      storage_modes,
        "top_complex":        top_complex,
        "workspace_names":    [ws.get("name", "") for ws in workspaces],
        # raw for table rendering
        "_workspaces": workspaces,
        "_datasets":   all_datasets,
        "_reports":    all_reports,
        "_measures":   all_measures,
    }


def _gen_fabric_executive_summary(client: openai.OpenAI, ctx: dict) -> str:
    complex_count = ctx["complexity_dist"].get("Very Complex", 0) + ctx["complexity_dist"].get("Complex", 0)
    import_count  = ctx["storage_modes"].get("Import", 0)
    prompt = f"""Write a 2-paragraph Executive Summary for a Microsoft Fabric Assessment Report.

Client: {ctx['client_name']}
Workspaces assessed: {ctx['workspace_count']} ({', '.join(ctx['workspace_names'][:5])})
Semantic Models (Datasets): {ctx['dataset_count']}
Reports: {ctx['report_count']}
Total Measures: {ctx['total_measures']}
Total Tables: {ctx['total_tables']}
Total Visuals: {ctx['total_visuals']}
Total Relationships: {ctx['total_relationships']}
Complex / Very Complex measures: {complex_count} of {ctx['total_measures']}
Import-mode datasets (refresh dependency): {import_count} of {ctx['dataset_count']}
Storage modes in use: {', '.join(f'{v} {k}' for k, v in ctx['storage_modes'].items())}

Paragraph 1: Describe the current Fabric/Power BI landscape — workspace structure, scale, model complexity.
Paragraph 2: Summarise key findings and modernisation recommendations (Direct Lake, Purview, governance).

Formal consultant tone. Cite specific numbers. No bullet points."""
    return _call_ai(client, prompt, max_tokens=600)


def _gen_fabric_ws_recommendations(client: openai.OpenAI, ctx: dict) -> str:
    complex_count = ctx["complexity_dist"].get("Very Complex", 0) + ctx["complexity_dist"].get("Complex", 0)
    import_count  = ctx["storage_modes"].get("Import", 0)
    prompt = f"""Write a 3-paragraph Fabric Modernisation Recommendations section.

Client: {ctx['client_name']}
Current state: {ctx['dataset_count']} semantic models, {ctx['report_count']} reports across {ctx['workspace_count']} workspaces.
{import_count} datasets use Import mode (scheduled refresh); {complex_count} measures are Complex or Very Complex.

Paragraph 1: Migrate Import-mode datasets to Direct Lake for real-time performance.
Paragraph 2: Rationalise complex DAX measures — refactor Very Complex measures, introduce reusable measure groups.
Paragraph 3: Governance — Microsoft Purview for lineage, sensitivity labels, workspace access tiers (Dev/UAT/Prod).

Formal consultant tone. Cite the numbers. No bullet points."""
    return _call_ai(client, prompt, max_tokens=700)


def _gen_fabric_conclusion(client: openai.OpenAI, ctx: dict) -> str:
    prompt = f"""Write a Conclusion paragraph (3–4 sentences) for a Fabric Assessment Report.
Client: {ctx['client_name']} — {ctx['dataset_count']} semantic models, {ctx['report_count']} reports,
{ctx['total_measures']} measures across {ctx['workspace_count']} workspaces.
Summarise the key findings and confirm the recommended next steps. Formal tone. 1 paragraph."""
    return _call_ai(client, prompt, max_tokens=250)


def build_fabric_ai_word_report(
    session_id: str,
    results: dict[str, Any],
    client_name: str | None = None,
) -> bytes:
    """
    Build an AI-powered Word report for a Fabric session.
    Generates Executive Summary, Recommendations, and Conclusion via GPT-4o.
    Data tables (workspace/model/report inventory, complexity) populated from results.
    """
    label    = client_name or f"Fabric Session {session_id[:8]}"
    run_date = datetime.utcnow().strftime("%d %b %Y")
    ctx      = _build_fabric_context(results, label)
    ai       = _get_client()

    logger.info("Fabric AI report: generating content for session %s", session_id)

    exec_summary      = _gen_fabric_executive_summary(ai, ctx)
    recommendations   = _gen_fabric_ws_recommendations(ai, ctx)
    conclusion        = _gen_fabric_conclusion(ai, ctx)

    logger.info("Fabric AI report: assembling Word document for session %s", session_id)

    doc = Document()
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    _setup_header(doc, label)
    _setup_footer(doc)
    _cover_page(doc, label, run_date)

    # ── Executive Summary ─────────────────────────────────────────────────────
    _h1(doc, "Executive Summary")
    for para in exec_summary.split("\n\n"):
        if para.strip():
            _body(doc, para.strip())

    # ── Workspace Overview ────────────────────────────────────────────────────
    _h1(doc, "Fabric Workspace Overview")
    ws_rows = []
    for ws in ctx["_workspaces"]:
        ds_list  = ws.get("datasets") or []
        rpt_list = ws.get("reports") or []
        measures = sum(len(d.get("measures") or []) for d in ds_list)
        ws_rows.append([
            ws.get("name", "—"),
            str(len(ds_list)),
            str(len(rpt_list)),
            str(measures),
            ws.get("state", "—"),
        ])
    _table(doc, ["Workspace", "Semantic Models", "Reports", "Total Measures", "State"],
           ws_rows or [["—", "—", "—", "—", "—"]], [5, 3, 2.5, 3.5, 2])

    # ── Semantic Model Inventory ──────────────────────────────────────────────
    _h1(doc, "Semantic Model Inventory")
    ds_rows = []
    for ds in ctx["_datasets"][:30]:
        complexity = ds.get("complexity_score", "—")
        ds_rows.append([
            ds.get("name", "—"),
            str(ds.get("table_count", "—")),
            str(ds.get("measure_count", "—")),
            str(ds.get("relationship_count", "—")),
            ds.get("storage_mode", "—"),
            f"{complexity}%" if isinstance(complexity, (int, float)) else str(complexity),
        ])
    _table(
        doc,
        ["Model Name", "Tables", "Measures", "Relationships", "Storage Mode", "Complexity"],
        ds_rows or [["—"] * 6],
        [5, 2, 2.5, 3, 3, 2.5],
    )

    # Storage mode summary
    _h2(doc, "Storage Mode Distribution")
    _table(
        doc,
        ["Storage Mode", "Count", "Implication"],
        [
            [mode, str(count),
             "Scheduled refresh required — Direct Lake migration recommended" if mode == "Import"
             else "Live query — consider Direct Lake for lakehouse sources" if mode == "DirectQuery"
             else "Optimal for Fabric OneLake — no refresh needed" if mode == "DirectLake"
             else "Mixed modes — review for consistency"]
            for mode, count in ctx["storage_modes"].items()
        ] or [["—", "—", "—"]],
        [3.5, 2, 10.5],
    )

    # ── Report Inventory ──────────────────────────────────────────────────────
    _h1(doc, "Report Inventory")
    rpt_rows = []
    for rp in ctx["_reports"][:30]:
        rpt_rows.append([
            rp.get("name", "—"),
            rp.get("report_type", "PowerBIReport"),
            str(rp.get("page_count") or "—"),
            str(rp.get("visual_count", "—")),
            str(rp.get("bookmark_count", "—")),
            "Yes" if rp.get("is_paginated") else "No",
        ])
    _table(
        doc,
        ["Report Name", "Type", "Pages", "Visuals", "Bookmarks", "Paginated"],
        rpt_rows or [["—"] * 6],
        [5.5, 3, 1.5, 2, 2.5, 2],
    )

    # ── DAX Complexity Analysis ───────────────────────────────────────────────
    _h1(doc, "DAX Complexity Analysis")

    _h2(doc, "Complexity Distribution")
    level_order = ["Very Complex", "Complex", "Moderate", "Simple", "None", "Unknown"]
    dist_rows = [
        [level, str(ctx["complexity_dist"].get(level, 0))]
        for level in level_order
        if ctx["complexity_dist"].get(level, 0) > 0
    ]
    _table(doc, ["Complexity Level", "Measure Count"],
           dist_rows or [["No measures analysed", "—"]], [8, 8])

    _h2(doc, "Top Complex Measures")
    top_rows = []
    for m in ctx["top_complex"]:
        cx = m.get("complexity") or {}
        top_rows.append([
            m.get("name", "—"),
            m.get("table", "—"),
            cx.get("level", "—"),
            str(cx.get("score", "—")),
            str(cx.get("function_count", "—")),
            str(cx.get("nesting_depth", "—")),
        ])
    _table(
        doc,
        ["Measure Name", "Table", "Level", "Score", "Functions", "Nesting Depth"],
        top_rows or [["No complex measures found", "—", "—", "—", "—", "—"]],
        [4.5, 3, 2.5, 1.5, 2, 2.5],
    )

    # ── Recommendations ───────────────────────────────────────────────────────
    _h1(doc, "Modernisation Recommendations")
    for para in recommendations.split("\n\n"):
        if para.strip():
            _body(doc, para.strip())

    doc.add_paragraph()
    _h2(doc, "Recommended Action Plan")
    _table(
        doc,
        ["Priority", "Action", "Benefit"],
        [
            ["High",   f"Migrate {ctx['storage_modes'].get('Import', 0)} Import-mode datasets to Direct Lake",
             "Eliminates scheduled refresh; near-real-time data for all reports"],
            ["High",   "Refactor Very Complex DAX measures",
             "Reduces query timeout risk and improves report load times"],
            ["Medium", "Implement workspace access tiers (Dev / UAT / Prod)",
             "Controlled deployment pipeline; reduces risk of breaking production reports"],
            ["Medium", "Enable Microsoft Purview sensitivity labels",
             f"Classifies sensitive columns across {ctx['dataset_count']} models"],
            ["Low",    "Consolidate duplicate semantic models across workspaces",
             "Reduces maintenance overhead and ensures single source of truth"],
        ],
        [2, 8, 6],
    )

    # ── Conclusion ────────────────────────────────────────────────────────────
    _h1(doc, "Conclusion")
    _h2(doc, "Summary of Findings")
    _table(
        doc,
        ["Metric", "Value"],
        [
            ["Workspaces Assessed",  str(ctx["workspace_count"])],
            ["Semantic Models",      str(ctx["dataset_count"])],
            ["Reports",              str(ctx["report_count"])],
            ["Total Measures",       str(ctx["total_measures"])],
            ["Total Tables",         str(ctx["total_tables"])],
            ["Total Visuals",        str(ctx["total_visuals"])],
            ["Total Relationships",  str(ctx["total_relationships"])],
            ["Complex Measures",     str(ctx["complexity_dist"].get("Complex", 0) + ctx["complexity_dist"].get("Very Complex", 0))],
            ["Import-mode Models",   str(ctx["storage_modes"].get("Import", 0))],
            ["Direct Lake Models",   str(ctx["storage_modes"].get("DirectLake", 0))],
        ],
        [8, 8],
    )
    _h2(doc, "Final Remarks")
    for para in conclusion.split("\n\n"):
        if para.strip():
            _body(doc, para.strip())

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    logger.info("Fabric AI report: Word document built for session %s", session_id)
    return buf.getvalue()
