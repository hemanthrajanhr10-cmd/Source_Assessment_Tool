"""Client Assessment Report builder.

Generates a professional .docx following the assessmentReport skill template:
  Cover → Executive Summary → Current State → Pain Points →
  Existing State Assessment → Inbound/Outbound → Fabric Recommendation →
  ROI & Licensing → Migration Estimation → Deliverable Summary
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from app.core.logging import get_logger
from app.models.client_assessment_report import ClientAssessmentReportRequest

logger = get_logger(__name__)

# ── Brand palette ─────────────────────────────────────────────────────────────
_DARK_BLUE  = RGBColor(0x1F, 0x38, 0x64)   # H1, cover title, table H3
_MID_BLUE   = RGBColor(0x2E, 0x75, 0xB6)   # H2, accent border
_TBL_HEADER = RGBColor(0x1F, 0x38, 0x64)   # Table header background
_TBL_ALT    = RGBColor(0xEB, 0xF3, 0xFB)   # Alternating row tint
_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
_MUTED      = RGBColor(0x59, 0x59, 0x59)
_BLACK      = RGBColor(0x00, 0x00, 0x00)


def _hex(c: RGBColor) -> str:
    return f"{c[0]:02X}{c[1]:02X}{c[2]:02X}"


# ── XML helpers ───────────────────────────────────────────────────────────────

def _shd(cell, colour: RGBColor) -> None:
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  _hex(colour))
    tcPr.append(shd)


def _cell_borders(tbl) -> None:
    for row in tbl.rows:
        for cell in row.cells:
            tc   = cell._tc
            tcPr = tc.get_or_add_tcPr()
            b    = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = OxmlElement(f"w:{side}")
                el.set(qn("w:val"),   "single")
                el.set(qn("w:sz"),    "4")
                el.set(qn("w:space"), "0")
                el.set(qn("w:color"), "CCCCCC")
                b.append(el)
            tcPr.append(b)


def _set_cell_margins(cell, top=60, bottom=60, left=100, right=100) -> None:
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar  = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"),    str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _run(para, text: str, size: float = 11, bold: bool = False,
         color: Optional[RGBColor] = None, font: str = "Arial") -> None:
    r = para.add_run(text)
    r.font.name  = font
    r.font.size  = Pt(size)
    r.font.bold  = bold
    if color:
        r.font.color.rgb = color


# ── Paragraph helpers ─────────────────────────────────────────────────────────

def _h1(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after  = Pt(6)
    _run(p, text, size=16, bold=True, color=_DARK_BLUE)
    # Blue underline accent
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot  = OxmlElement("w:bottom")
    bot.set(qn("w:val"),   "single")
    bot.set(qn("w:sz"),    "6")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), _hex(_MID_BLUE))
    pBdr.append(bot)
    pPr.append(pBdr)


def _h2(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    _run(p, text, size=13, bold=True, color=_MID_BLUE)


def _h3(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(3)
    _run(p, text, size=12, bold=True, color=_DARK_BLUE)


def _body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    _run(p, text, size=11)


def _bullet(doc: Document, text: str, bold_label: str = "") -> None:
    try:
        p = doc.add_paragraph(style="List Bullet")
    except Exception:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    if bold_label:
        _run(p, f"{bold_label} — ", size=11, bold=True)
    _run(p, text, size=11)


def _spacer(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)


# ── Table helper ──────────────────────────────────────────────────────────────

def _table(doc: Document, headers: List[str], rows: List[List[str]]) -> None:
    if not rows:
        return
    n_cols = len(headers)
    tbl = doc.add_table(rows=1 + len(rows), cols=n_cols)
    tbl.style = "Table Grid"

    # Header row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        cell = hdr_cells[i]
        _shd(cell, _TBL_HEADER)
        _set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.clear()
        _run(p, h, size=10, bold=True, color=_WHITE)

    # Data rows
    for ri, row_data in enumerate(rows):
        tr_cells = tbl.rows[ri + 1].cells
        shade = ri % 2 == 1
        for ci, val in enumerate(row_data):
            cell = tr_cells[ci]
            if shade:
                _shd(cell, _TBL_ALT)
            _set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.clear()
            _run(p, str(val) if val is not None else "", size=10)

    _cell_borders(tbl)
    _spacer(doc)


# ── Main builder ──────────────────────────────────────────────────────────────

def build_client_assessment_report(data: ClientAssessmentReportRequest) -> bytes:
    doc = Document()

    # Page setup
    for sec in doc.sections:
        sec.page_width    = Inches(8.5)
        sec.page_height   = Inches(11)
        sec.top_margin    = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin   = Cm(2.5)
        sec.right_margin  = Cm(2.5)

    arch = data.architecture_description or "traditional BI architecture"
    src_names = ", ".join(s.name for s in data.data_sources) if data.data_sources else "multiple source systems"

    # ── 1. Cover Page ─────────────────────────────────────────────────────────
    for _ in range(4):
        doc.add_paragraph()

    cover_title = doc.add_paragraph()
    cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(cover_title, "Assessment Report", size=28, bold=True, color=_DARK_BLUE)
    cover_title.paragraph_format.space_after = Pt(12)

    cover_client = doc.add_paragraph()
    cover_client.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(cover_client, data.client_name, size=36, bold=True, color=_MID_BLUE)
    cover_client.paragraph_format.space_after = Pt(12)

    cover_sub = doc.add_paragraph()
    cover_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(cover_sub, "Microsoft Fabric Modernization Assessment", size=14, color=_MUTED)

    cover_date = doc.add_paragraph()
    cover_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(cover_date, datetime.now().strftime("%B %Y"), size=12, color=_MUTED)

    doc.add_page_break()

    # ── 2. Executive Summary ──────────────────────────────────────────────────
    _h1(doc, "Executive Summary")
    _body(doc,
        f"{data.client_name} currently operates a {arch} serving approximately "
        f"{data.user_count} report consumers and {data.dev_count} report developers. "
        f"While the environment supports current operational needs, it faces significant "
        f"challenges that limit scalability, agility, and innovation in analytics."
    )
    _body(doc,
        f"The existing environment faces several key challenges including data silos, legacy "
        f"ETL processes, limited real-time capabilities, governance gaps, and the absence of "
        f"proper development environments. These constraints restrict {data.client_name}'s "
        f"ability to deliver timely, trusted insights to stakeholders."
    )
    _body(doc,
        f"To address these challenges, {data.client_name} is seeking to modernize its data "
        f"platform. Microsoft Fabric presents a unified, scalable solution that can:"
    )
    benefits = data.fabric_benefits or [
        "Eliminate data silos through a centralized data platform and Single Source of Truth",
        "Enable real-time analytics and streaming capabilities",
        "Support advanced AI use cases and Copilot integration",
        "Provide robust governance through integrated tooling",
        "Establish proper Dev/Test/Prod environments for controlled deployments",
    ]
    for b in benefits:
        _bullet(doc, b)
    _body(doc,
        f"This transformation will position {data.client_name} to move from a legacy, tightly "
        f"coupled architecture to a modern, flexible, and AI-enabled data platform."
    )

    # ── 3. Current State Overview ─────────────────────────────────────────────
    doc.add_page_break()
    _h1(doc, "Current State Overview")
    _body(doc,
        f"{data.client_name} currently operates a {arch} built on {src_names}."
    )
    _body(doc, "The primary data flow is structured as follows:")
    flow = data.data_flow_steps or [
        "Data from primary sources is ingested into staging databases",
        "Data is moved and transformed using ETL processes",
        "Processed data is stored in the reporting/warehouse database",
        "A semantic model is built on top of this layer serving reports",
        "Reports are consumed via Power BI and/or Excel",
    ]
    for s in flow:
        _bullet(doc, s)
    _body(doc, "The environment supports approximately:")
    _bullet(doc, f"{data.user_count} report consumers")
    _bullet(doc, f"{data.dev_count} report developers, responsible for report creation and maintenance")

    # ── 4. Key Pain Points ────────────────────────────────────────────────────
    doc.add_page_break()
    _h1(doc, "Key Pain Points")
    for pp in data.pain_points:
        _bullet(doc, pp.description, pp.label)

    # ── 5. Existing State Assessment ──────────────────────────────────────────
    doc.add_page_break()
    _h1(doc, "Existing State Assessment")

    _h2(doc, "Existing Architecture")
    _body(doc, f"The current architecture at {data.client_name} relies on multiple interconnected systems.")

    _h2(doc, "Data Sources Overview")
    _body(doc, "The following data sources feed the current BI platform:")
    _table(doc,
        ["Category", "System Name", "Type", "Integration Method", "Description"],
        [[s.category, s.name, s.type, s.integration_method, s.description]
         for s in data.data_sources],
    )

    for db in data.databases:
        _h3(doc, db.name)
        _table(doc,
            ["Database Name", "Tables", "Views", "Stored Procs", "Functions", "Data Size"],
            [[db.name, db.table_count, db.view_count, db.stored_proc_count, db.function_count, db.data_size]],
        )
        for obs in db.observations:
            _bullet(doc, obs)
        if db.high_volume_tables:
            _h3(doc, f"High-Volume Tables — {db.name}")
            _table(doc,
                ["Schema", "Table Name", "Row Count", "Activity Level"],
                [[t.schema_name, t.name, t.row_count, t.activity] for t in db.high_volume_tables],
            )

    _h2(doc, "ETL / Orchestration Details")
    job_note = (
        f"Based on the assessment, {data.etl_job_count} jobs are configured."
        if data.etl_job_count else ""
    )
    _body(doc,
        f"The current environment uses {data.etl_tool} as the primary orchestration mechanism. "
        f"{job_note}"
    )
    if data.etl_jobs:
        _table(doc,
            ["S.No", "Job / Package Name"],
            [[str(i + 1), j] for i, j in enumerate(data.etl_jobs)],
        )

    _h2(doc, "Reporting Layer Overview")
    _body(doc,
        f"The reporting layer at {data.client_name} is built on {data.reporting_tool}, "
        f"serving {data.user_count} users."
    )
    if data.semantic_model:
        _h3(doc, "Semantic Model — Structure Overview")
        _table(doc,
            ["Component", "Count"],
            [
                ["Tables",        data.semantic_model.table_count],
                ["Measures",      data.semantic_model.measure_count],
                ["Relationships", data.semantic_model.relationship_count],
                ["Roles",         data.semantic_model.role_count],
            ],
        )
    if data.reports:
        _h3(doc, "Power BI Reporting Landscape")
        _table(doc,
            ["Platform", "Workspace", "Report Name", "Description"],
            [[r.platform, r.workspace, r.name, r.description] for r in data.reports],
        )

    # ── 6. Inbound & Outbound Systems ─────────────────────────────────────────
    doc.add_page_break()
    _h1(doc, "Inbound and Outbound Systems Overview")

    _h2(doc, "Inbound Systems (Data Sources to BI Platform)")
    _table(doc,
        ["Category", "System Name", "Type", "Integration Method", "Description"],
        [[s.category, s.name, s.type, s.integration_method, s.description]
         for s in data.data_sources],
    )

    _h2(doc, "Outbound Systems (BI Consumption Layer)")
    outbound_rows = (
        [[s.category, s.name, s.type, s.connection, s.description] for s in data.outbound_systems]
        if data.outbound_systems else
        [["BI Reporting", "Power BI", "Visualization Tool", "Live Connection / Import",
          f"Primary reporting platform (~{data.user_count} users)"]]
    )
    _table(doc,
        ["Category", "System Name", "Type", "Connection Type", "Description"],
        outbound_rows,
    )

    _h2(doc, "Intermediate Systems (Processing Layer)")
    if data.intermediate_systems:
        _table(doc,
            ["Layer", "System Name", "Role"],
            [[s.layer, s.name, s.role] for s in data.intermediate_systems],
        )

    # ── 7. Fabric Recommendation ──────────────────────────────────────────────
    doc.add_page_break()
    _h1(doc, "Fabric Recommendation")

    _h2(doc, "Proposed Architecture")
    _body(doc,
        "The proposed Microsoft Fabric architecture replaces multiple legacy components with a "
        "unified, Medallion-based platform (Bronze → Silver → Gold), eliminating data silos and "
        "enabling real-time and AI-driven analytics."
    )

    for rec in data.source_recommendations:
        _h3(doc, rec.source_name)
        for r in rec.recommendations:
            _bullet(doc, r)

    _h2(doc, "Azure DevOps — Code Check-in Strategy")
    _body(doc,
        "It is recommended to implement a Git-integrated DevOps approach in Microsoft Fabric "
        "using separate Dev and Prod workspaces."
    )
    for b in [
        "Dev Workspace — Development & testing (Git-integrated)",
        "Prod Workspace — Production (restricted access, governed deployments)",
        "Branching Strategy: feature/* → dev → main",
        "Mandatory Pull Request (PR) with peer review before merging",
        "All Fabric artifacts stored in Git — pipelines, notebooks, semantic models",
        "Use Fabric Deployment Pipelines to promote Dev → Prod in a controlled manner",
    ]:
        _bullet(doc, b)

    _h2(doc, "Monitoring")
    for b in [
        "Fabric Monitor Hub tracks all pipeline runs and notebook executions in real time",
        "Data Activator configured for alerts (pipeline failures, refresh SLA breaches, capacity throttling)",
        "All pipelines configured with automatic retries (3 attempts) before failure alert",
        "Centralized audit trail stored in OneLake for full pipeline visibility",
    ]:
        _bullet(doc, b)

    _h2(doc, "Power BI Report Creation")
    for b in [
        "Migrate reporting to Fabric using a Gold Warehouse Direct Lake semantic model",
        "Consolidate into a single reusable semantic model — eliminate duplicate datasets",
        "Follow controlled Dev → Prod deployment via Deployment Pipeline",
        "Migrate and enforce Row-Level Security (RLS) — preserve all existing role-based access",
        "Store all report and model definitions in Azure DevOps for version control",
    ]:
        _bullet(doc, b)

    _h2(doc, "Real-Time Intelligence")
    for b in [
        "Use database mirroring or incremental ingestion to capture near-live transactional data",
        "Build dashboards for live order pipeline, real-time inventory, and intraday financial activity",
        "Combine real-time data with historical Gold layer data for hybrid reporting",
        "Use Data Activator for event-driven alerts — inventory thresholds, sales underperformance, anomalies",
    ]:
        _bullet(doc, b)

    _h2(doc, "AI Enablement — Data Agent")
    for b in [
        "Use Fabric Data Agent (Copilot) on top of the Gold Warehouse for natural language queries",
        "Leverage existing business logic, KPIs, and measures from Power BI reports to configure the Data Agent",
        "Apply Role-Based Access Control (RBAC) to restrict data visibility based on user roles",
        "Publish the Data Agent to Microsoft Teams via Copilot Studio for everyday use",
    ]:
        _bullet(doc, b)

    # ── 8. ROI & Licensing ────────────────────────────────────────────────────
    doc.add_page_break()
    _h1(doc, "ROI & Licensing")

    _h2(doc, "Cost Comparison")
    _body(doc,
        f"The current {data.client_name} environment uses multiple separate technologies, creating "
        f"recurring costs across storage, ETL, semantic modelling, and reporting. The proposed "
        f"Microsoft Fabric architecture consolidates these into a single capacity-based platform."
    )
    if data.cost_comparison:
        _table(doc,
            ["Category", "Existing Setup", f"Microsoft Fabric ({data.proposed_capacity} Capacity)"],
            [[c.category, c.current_cost, c.proposed_cost] for c in data.cost_comparison],
        )

    _h2(doc, "ROI (Return on Investment)")
    roi_rows = (
        [[r.category, r.current, r.proposed, r.benefit] for r in data.roi_table]
        if data.roi_table else [
            ["Cost Efficiency",          "Existing platform cost",    "Reduced Fabric cost",       "10-30% Savings"],
            ["Platforms Managed",        "5+ Separate Tools",         "Single Unified Platform",   "80% Simpler Architecture"],
            ["Operational Productivity", "High manual support",       "Managed SaaS Platform",     "30%+ Lower Effort"],
            ["Scalability",              "Multi-system expansion",    "Single Capacity Scale",     "50% Faster Delivery"],
            ["Governance",               "Limited lineage",           "Centralized governance",    "Improved compliance"],
        ]
    )
    _table(doc, ["Category", "Existing Setup", "Microsoft Fabric", "ROI Benefit"], roi_rows)

    _h2(doc, "Licensing")
    _body(doc,
        f"Based on current workload and future scalability requirements, {data.proposed_capacity} "
        f"capacity is recommended for {data.client_name}."
    )
    _bullet(doc, f"Region: {data.region}")
    if data.reserved_price:
        _bullet(doc, f"1-Year Reservation: {data.reserved_price}/month (recommended)")
    if data.paygo_price:
        _bullet(doc, f"Pay-as-you-go: {data.paygo_price}/month")
    _bullet(doc, "OneLake Storage: $0.023 per GB/month")
    _bullet(doc, "Reserved capacity offers approximately 40% cost savings compared to PAYG")

    # ── 9. Migration Estimation ───────────────────────────────────────────────
    doc.add_page_break()
    _h1(doc, "Migration Estimation")
    _body(doc,
        "The implementation is estimated based on the current scope of data sources, "
        "transformations, and reporting requirements."
    )
    _body(doc, f"Project Duration: {data.project_duration}")

    task_rows = (
        [[t.task, t.hours] for t in data.migration_tasks]
        if data.migration_tasks else [
            ["Source & Report Analysis",             "40"],
            ["Data Ingestion Setup (all sources)",   "80"],
            ["Bronze Layer Setup",                   "16"],
            ["Silver Layer Setup",                   "80"],
            ["Gold Layer Setup (Business Logic)",    "160"],
            ["DevOps & Git Integration",             "40"],
            ["Semantic Model Development",           "120"],
            ["Power BI Report Migration",            "200"],
            ["AI / Data Agent Development",          "80"],
            ["Monitoring & Alerts Setup",            "16"],
            ["Testing & Validation",                 "80"],
            ["Deployment & Release Management",      "24"],
            ["Post Production Support",              "40"],
            ["Project Management",                   "160"],
            ["Total Hours",                          "1136"],
        ]
    )
    _table(doc, ["Tasks", "Hours"], task_rows)

    _h2(doc, "Resource Split-up")
    resource_rows = (
        [[r.resource, r.hours] for r in data.resource_split]
        if data.resource_split else [
            ["Project Manager (1)",         "160"],
            ["QA (1)",                      "120"],
            ["Senior Engineer (1)",         "320"],
            ["Junior / Mid Engineer (2)",   "536"],
            ["Total",                       "1136"],
        ]
    )
    _table(doc, ["Resources", "Hours"], resource_rows)

    # ── 10. Deliverable Summary ───────────────────────────────────────────────
    doc.add_page_break()
    _h1(doc, "Deliverable Summary")
    _body(doc,
        f"The proposed implementation will transform {data.client_name}'s current reporting "
        f"platform into a modern, scalable, and governed analytics environment. The following "
        f"key deliverables will be implemented as part of this engagement:"
    )
    deliverables = data.deliverables or [
        "Medallion Architecture with Bronze (raw), Silver (cleansed), and Gold (business-ready) layers",
        "Near real-time data ingestion from all primary source systems",
        "Migration of all ETL logic from legacy processes into Fabric Pipelines and Notebooks",
        "Reimplementation of SQL views, stored procedures, and business logic in Fabric Gold Warehouse",
        "Centralized Fabric Semantic Model replacing the existing semantic/cube layer",
        "Migration and modernization of all Power BI reports into governed Fabric-based models",
        "Setup of Dev and Production workspaces with Git integration and deployment pipelines",
        "Data Agent implementation for natural language querying over the Gold layer",
        "Monitoring, alerting, logging, and audit tracking using Fabric-native tools",
        "Improved governance through centralized security, access control, and lineage",
    ]
    for d in deliverables:
        _bullet(doc, d)

    # ── Serialize ─────────────────────────────────────────────────────────────
    buf = io.BytesIO()
    doc.save(buf)
    logger.info("Client assessment report generated for %s", data.client_name)
    return buf.getvalue()
