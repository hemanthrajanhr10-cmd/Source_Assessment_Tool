"""
One-off script: pull Fabric assessment data from Azure SQL and generate Excel.
"""
import json
from collections import defaultdict

import mssql_python
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

USER_ID  = "9b08f27f-27a6-4b0f-aec2-3075e29ddc40"
CONN_STR = (
    "SERVER=tcp:uiap-source-assessment.database.windows.net,1433;"
    "DATABASE=SourceAssessment;UID=sqladmin;PWD=welcome@1234;"
    "TrustServerCertificate=no;Encrypt=yes;"
)
OUT_PATH = r"C:\Users\hemanth.rajan\Downloads\Farnell_Fabric_Assessment_Summary.xlsx"


def fetch_results():
    conn = mssql_python.connect(CONN_STR)
    cur  = conn.cursor()
    cur.execute(
        "SELECT results_json FROM dbo.fabric_sessions "
        "WHERE user_id = ? ORDER BY created_at DESC",
        (USER_ID,),
    )
    row = cur.fetchone()
    conn.close()
    if row is None:
        raise RuntimeError("No session found.")
    return json.loads(row[0])


def risk_level(score, rc):
    if score >= 40 or rc >= 5:
        return "High"
    if score >= 20 or rc >= 2:
        return "Medium"
    return "Low"


def complexity_level(score):
    if score >= 40: return "Very High"
    if score >= 25: return "High"
    if score >= 15: return "Medium"
    if score >= 5:  return "Low"
    return "Very Low"


def build_data(data):
    workspaces = data.get("workspaces", [])
    ds_map = {}
    for ws in workspaces:
        for ds in ws.get("semantic_models", ws.get("datasets", [])):
            ds_map[ds["id"]] = {
                "name":              ds["name"],
                "workspace":         ws["name"],
                "table_count":       ds.get("table_count", 0),
                "measure_count":     ds.get("measure_count", 0),
                "relationship_count":ds.get("relationship_count", 0),
                "complexity_score":  ds.get("complexity_score", 0),
                "storage_mode":      ds.get("storage_mode") or "Import",
            }

    ds_reports = defaultdict(list)
    for ws in workspaces:
        for rp in ws.get("reports", []):
            dsid = rp.get("dataset_id", "")
            if dsid:
                ds_reports[dsid].append(rp["name"])

    db_rows = []
    for dsid, ds in ds_map.items():
        rc = len(ds_reports.get(dsid, []))
        sc = ds["complexity_score"]
        dep = "Shared Dataset" if rc > 1 else ("Standalone" if rc == 1 else "Unused")
        db_rows.append({
            "Database (Semantic Model)":    ds["name"],
            "Workspace":                    ds["workspace"],
            "Used By (Approx Report Count)":rc,
            "Dependency Type":              dep,
            "Risk Level":                   risk_level(sc, rc),
        })
    db_rows.sort(key=lambda x: -x["Used By (Approx Report Count)"])

    report_rows = []
    for ws in workspaces:
        for rp in ws.get("reports", []):
            dsid  = rp.get("dataset_id", "")
            ds    = ds_map.get(dsid, {})
            score = ds.get("complexity_score", 0)
            report_rows.append({
                "Report Name":      rp["name"],
                "Workspace":        ws["name"],
                "Table Count":      ds.get("table_count", 0),
                "Measures":         ds.get("measure_count", 0),
                "Relationships":    ds.get("relationship_count", 0),
                "Data Sources":     ds.get("storage_mode") or "Import",
                "Complexity Score": score,
                "Complexity Level": complexity_level(score),
            })
    report_rows.sort(key=lambda x: -x["Complexity Score"])
    return db_rows, report_rows


def make_workbook(db_rows, report_rows):
    wb = openpyxl.Workbook()
    HDR_FILL   = PatternFill("solid", fgColor="1F3864")
    HDR_FONT   = Font(bold=True, color="FFFFFF", size=11)
    TITLE_FILL = PatternFill("solid", fgColor="2E75B6")
    TITLE_FONT = Font(bold=True, color="FFFFFF", size=13)
    ODD_FILL   = PatternFill("solid", fgColor="EBF3FB")
    EVEN_FILL  = PatternFill("solid", fgColor="FFFFFF")
    RISK_CLRS  = {"High":"C00000","Medium":"ED7D31","Low":"70AD47"}
    CMPLX_CLRS = {"Very High":"C00000","High":"FF0000","Medium":"FFC000","Low":"70AD47","Very Low":"00B050"}

    thin = Side(style="thin", color="BDD7EE")
    bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)
    cnt  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    lft  = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    def add_title(wsxl, text, n):
        wsxl.merge_cells(f"A1:{get_column_letter(n)}1")
        c = wsxl["A1"]
        c.value=text; c.fill=TITLE_FILL; c.font=TITLE_FONT; c.alignment=cnt
        wsxl.row_dimensions[1].height=28

    def add_header(wsxl, row, cols):
        for ci,val in enumerate(cols,1):
            c=wsxl.cell(row=row,column=ci,value=val)
            c.fill=HDR_FILL; c.font=HDR_FONT; c.alignment=cnt; c.border=bdr
        wsxl.row_dimensions[row].height=22

    def add_data_row(wsxl, row, vals, cols, is_odd):
        fill=ODD_FILL if is_odd else EVEN_FILL
        for ci,key in enumerate(cols,1):
            c=wsxl.cell(row=row,column=ci,value=vals[key])
            c.fill=fill; c.alignment=lft; c.border=bdr
        wsxl.row_dimensions[row].height=18

    # Sheet 1
    ws1=wb.active; ws1.title="Database Usage Summary"
    cols1=["Database (Semantic Model)","Workspace","Used By (Approx Report Count)","Dependency Type","Risk Level"]
    add_title(ws1,"Database Usage Summary",len(cols1))
    add_header(ws1,2,cols1)
    for ri,rd in enumerate(db_rows,3):
        add_data_row(ws1,ri,rd,cols1,ri%2==1)
        rl=rd["Risk Level"]
        c=ws1.cell(row=ri,column=5)
        c.fill=PatternFill("solid",fgColor=RISK_CLRS.get(rl,"FFFFFF"))
        c.font=Font(bold=True,color="FFFFFF",size=10); c.alignment=cnt
    for i,w in enumerate([38,18,30,22,14],1):
        ws1.column_dimensions[get_column_letter(i)].width=w

    # Sheet 2
    ws2=wb.create_sheet("Report Complexity")
    cols2=["Report Name","Workspace","Table Count","Measures","Relationships","Data Sources","Complexity Score","Complexity Level"]
    add_title(ws2,"Detailed Report Complexity",len(cols2))
    add_header(ws2,2,cols2)
    for ri,rd in enumerate(report_rows,3):
        add_data_row(ws2,ri,rd,cols2,ri%2==1)
        cl=rd["Complexity Level"]
        c=ws2.cell(row=ri,column=8)
        c.fill=PatternFill("solid",fgColor=CMPLX_CLRS.get(cl,"FFFFFF"))
        c.font=Font(bold=True,color="FFFFFF" if cl in("Very High","High","Medium") else "000000",size=10)
        c.alignment=cnt
        for ci in [3,4,5,7]:
            ws2.cell(row=ri,column=ci).alignment=cnt
    for i,w in enumerate([32,16,14,12,16,18,18,18],1):
        ws2.column_dimensions[get_column_letter(i)].width=w
    return wb


def main():
    print("Fetching data from Azure SQL...")
    data=fetch_results()
    db_rows,report_rows=build_data(data)
    print(f"  Datasets : {len(db_rows)}")
    print(f"  Reports  : {len(report_rows)}")
    wb=make_workbook(db_rows,report_rows)
    wb.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")

if __name__=="__main__":
    main()
